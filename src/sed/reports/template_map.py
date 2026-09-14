"""PowerPoint template maps: which layout, placeholders and content box each slide kind uses in a given template.

A map is YAML (`templates/pptx/<name>.map.yaml` in the repo, or `DATA_DIR/config/templates/<name>.map.yaml` on a
profile). The corporate template and its map live only in DATA_DIR; the repo ships the neutral map for python-pptx's
built-in default template. Changing templates is a map change, never a code change.

Lookup order for `load_template_map(name_or_path, paths)`: an explicit path, then DATA_DIR, then the repo; `None`
means `settings.reports.template_map`. Layouts resolve by name, then by `fallback_index`. Any mismatch between map and
template (missing layout or placeholder, slide size, content box outside the slide) is a validation error (exit 2):
the builder never resizes a template.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal, get_args

from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator, model_validator

from sed.errors import ValidationFailed
from sed.paths import Paths, repo_root
from sed.reports.specs import SlideKind

SLIDE_KINDS: tuple[str, ...] = get_args(SlideKind)
TEXT_KINDS = frozenset({"title", "section", "narrative"})  # kinds that may omit content_box_in
EMU_PER_INCH = 914400
SLIDE_RATIOS = {"4:3": 4 / 3, "16:9": 16 / 9}
RATIO_TOLERANCE = 0.01
MAP_SUFFIX = ".map.yaml"
POTX_MESSAGE = "python-pptx cannot open .potx templates: open it in PowerPoint and save as .pptx"
_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,63}$")
_HEX_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class FontSpec(_Strict):
    title: str = Field(min_length=1)
    body: str = Field(min_length=1)
    body_pt: float = Field(ge=6, le=40)
    table_pt: float = Field(ge=6, le=28)


class LayoutSpec(_Strict):
    layout_name: str = Field(min_length=1)
    fallback_index: int = Field(ge=0)
    placeholders: dict[Literal["title", "subtitle", "body"], int] = Field(default_factory=dict)
    content_box_in: list[float] | None = None
    max_rows: int = Field(12, ge=1, le=60)
    max_bullets: int = Field(6, ge=1, le=20)

    @field_validator("placeholders")
    @classmethod
    def _idx_non_negative(cls, v: dict[str, int]) -> dict[str, int]:
        bad = [k for k, idx in v.items() if idx < 0]
        if bad:
            raise ValueError(f"placeholder idx must be >= 0 ({', '.join(bad)})")
        return v

    @field_validator("content_box_in")
    @classmethod
    def _box(cls, v: list[float] | None) -> list[float] | None:
        if v is None:
            return v
        if len(v) != 4:
            raise ValueError("content_box_in is [x, y, width, height] in inches")
        x, y, w, h = v
        if x < 0 or y < 0 or w <= 0 or h <= 0:
            raise ValueError("content_box_in needs x, y >= 0 and width, height > 0")
        return v


class TemplateMap(_Strict):
    name: str
    template: str | None = None
    slide_size: Literal["4:3", "16:9"]
    fonts: FontSpec
    series_colors: list[str] = Field(min_length=3)
    footer_text: str = "{report_title} | {period}"
    classification_label: str = ""
    layouts: dict[SlideKind, LayoutSpec]

    @field_validator("name")
    @classmethod
    def _name(cls, v: str) -> str:
        if not _NAME_RE.match(v):
            raise ValueError("name may only use letters, digits, '_', '.' and '-'")
        return v

    @field_validator("series_colors")
    @classmethod
    def _colors(cls, v: list[str]) -> list[str]:
        bad = [c for c in v if not _HEX_RE.match(c)]
        if bad:
            raise ValueError(f"series_colors must be '#RRGGBB' hex colours (got {bad})")
        return v

    @model_validator(mode="after")
    def _all_kinds(self) -> TemplateMap:
        missing = [k for k in SLIDE_KINDS if k not in self.layouts]
        if missing:
            raise ValueError(f"layouts must define every slide kind; missing: {', '.join(missing)}")
        no_box = [k for k, spec in self.layouts.items() if k not in TEXT_KINDS and spec.content_box_in is None]
        if no_box:
            raise ValueError(f"content_box_in is required for: {', '.join(no_box)}")
        for kind in ("title", "section", "narrative"):
            if "title" not in self.layouts[kind].placeholders:
                raise ValueError(f"layouts.{kind} needs a title placeholder")
        if "body" not in self.layouts["narrative"].placeholders and self.layouts["narrative"].content_box_in is None:
            raise ValueError("layouts.narrative needs a body placeholder or content_box_in")
        return self


@dataclass(frozen=True)
class LoadedTemplateMap:
    map: TemplateMap
    path: Path
    template_path: Path | None  # None: python-pptx's built-in default template
    sha256: str


def default_template_path() -> Path:
    """python-pptx's built-in default template (4:3, no branding)."""
    import pptx

    return Path(pptx.__file__).resolve().parent / "templates" / "default.pptx"


def _looks_like_path(value: str) -> bool:
    lower = value.lower()
    return "/" in value or "\\" in value or lower.endswith((".yaml", ".yml", ".pptx", ".potx"))


def find_map_file(name_or_path: str, paths: Paths | None) -> Path:
    """Resolve a map name or path to a file (explicit path, DATA_DIR/config/templates, repo templates/pptx)."""
    value = name_or_path.strip()
    if value.lower().endswith(".potx"):
        raise ValidationFailed(POTX_MESSAGE, {"file": value})
    if value.lower().endswith(".pptx"):
        raise ValidationFailed(
            f"'{value}' is a template, not a template map; pass a {MAP_SUFFIX} file that names it in 'template:'"
        )
    if _looks_like_path(value):
        path = Path(value)
        if not path.is_file():
            raise ValidationFailed(f"Template map file not found: {path}")
        return path.resolve()
    if not _NAME_RE.match(value):
        raise ValidationFailed(f"Invalid template map name '{value}'")
    candidates = []
    if paths is not None:
        candidates.append(paths.config / "templates" / f"{value}{MAP_SUFFIX}")
    candidates.append(repo_root() / "templates" / "pptx" / f"{value}{MAP_SUFFIX}")
    for candidate in candidates:
        if candidate.is_file():
            return candidate.resolve()
    raise ValidationFailed(f"Template map '{value}' not found", {"searched": [str(c) for c in candidates]})


def _parse_map(path: Path) -> tuple[TemplateMap, bytes]:
    from sed.settings import read_yaml

    raw = path.read_bytes()
    data = read_yaml(path)
    try:
        return TemplateMap.model_validate(data), raw
    except ValidationError as exc:
        errors = [{"loc": ".".join(str(p) for p in e["loc"]), "msg": e["msg"]} for e in exc.errors()]
        raise ValidationFailed(f"Invalid template map {path}", errors) from exc


def _template_file(tmap: TemplateMap, map_path: Path) -> Path | None:
    if not tmap.template:
        return None
    template = Path(tmap.template)
    if not template.is_absolute():
        template = map_path.parent / template
    if template.suffix.lower() == ".potx":
        raise ValidationFailed(POTX_MESSAGE, {"file": str(template)})
    if template.suffix.lower() != ".pptx":
        raise ValidationFailed(f"Template must be a .pptx file (got {template.name})")
    if not template.is_file():
        raise ValidationFailed(f"Template file not found: {template}", {"map": str(map_path)})
    return template.resolve()


def load_template_map(name_or_path: str | None = None, paths: Paths | None = None) -> LoadedTemplateMap:
    """Load, validate and hash a template map; the template is opened to check layouts, placeholders and size."""
    if name_or_path is None:
        from sed.settings import load_settings

        name_or_path = load_settings(paths).reports.template_map
    map_path = find_map_file(name_or_path, paths)
    tmap, raw = _parse_map(map_path)
    template = _template_file(tmap, map_path)
    template_bytes = (template or default_template_path()).read_bytes()
    digest = hashlib.sha256()
    digest.update(raw)
    digest.update(b"\0")
    digest.update(template_bytes)
    loaded = LoadedTemplateMap(tmap, map_path, template, digest.hexdigest())
    check_template(open_template(loaded, drop_slides=False), tmap)
    return loaded


def open_template(loaded: LoadedTemplateMap, *, drop_slides: bool = True) -> Any:
    """Open the map's template (python-pptx default when `template` is null), optionally without its sample slides."""
    prs = open_pptx(loaded.template_path)
    if drop_slides:
        slide_ids = prs.slides._sldIdLst
        for slide_id in list(slide_ids):
            prs.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
    return prs


def open_pptx(path: Path | None) -> Any:
    """Open a .pptx (exit 2 for .potx, missing or unreadable files)."""
    from pptx import Presentation

    if path is None:
        return Presentation()
    if path.suffix.lower() == ".potx":
        raise ValidationFailed(POTX_MESSAGE, {"file": str(path)})
    if not path.is_file():
        raise ValidationFailed(f"Template file not found: {path}")
    try:
        return Presentation(str(path))
    except Exception as exc:  # python-pptx raises several unrelated types for non-pptx input
        raise ValidationFailed(f"Cannot open {path.name} as a PowerPoint .pptx file", {"error": str(exc)}) from exc


def slide_ratio_name(width_emu: int, height_emu: int) -> str:
    ratio = width_emu / height_emu if height_emu else 0.0
    for name, expected in SLIDE_RATIOS.items():
        if abs(ratio - expected) / expected <= RATIO_TOLERANCE:
            return name
    return "other"


def all_layouts(prs: Any) -> list[tuple[int, int, Any]]:
    """(master index, layout index within its master, layout) for every layout of every master."""
    out = []
    for m_idx, master in enumerate(prs.slide_masters):
        out += [(m_idx, l_idx, layout) for l_idx, layout in enumerate(master.slide_layouts)]
    return out


def resolve_layout(prs: Any, spec: LayoutSpec) -> tuple[Any, bool]:
    """The layout named `layout_name` (any master), else the first master's layout at `fallback_index`.

    Returns (layout, fallback_used); raises exit 2 when neither exists.
    """
    for _, _, layout in all_layouts(prs):
        if layout.name == spec.layout_name:
            return layout, False
    layouts = prs.slide_layouts
    if spec.fallback_index < len(layouts):
        return layouts[spec.fallback_index], True
    raise ValidationFailed(
        f"Layout '{spec.layout_name}' not found and fallback_index {spec.fallback_index} is out of range",
        {"available": [layout.name for _, _, layout in all_layouts(prs)]},
    )


def check_template(prs: Any, tmap: TemplateMap) -> None:
    """Validate that a template satisfies a map: size ratio, layouts, declared placeholders and content boxes."""
    actual = slide_ratio_name(prs.slide_width, prs.slide_height)
    if actual != tmap.slide_size:
        raise ValidationFailed(
            f"Template slide size is {actual} but the map declares {tmap.slide_size}; the builder never resizes",
            {
                "width_in": round(prs.slide_width / EMU_PER_INCH, 3),
                "height_in": round(prs.slide_height / EMU_PER_INCH, 3),
            },
        )
    width_in = prs.slide_width / EMU_PER_INCH
    height_in = prs.slide_height / EMU_PER_INCH
    errors: list[dict[str, Any]] = []
    for kind in SLIDE_KINDS:
        spec = tmap.layouts[kind]
        layout, _ = resolve_layout(prs, spec)
        present = {ph.placeholder_format.idx for ph in layout.placeholders}
        for role, idx in spec.placeholders.items():
            if idx not in present:
                errors.append(
                    {
                        "kind": kind,
                        "msg": f"layout '{layout.name}' has no placeholder idx {idx} for {role}",
                        "available": sorted(present),
                    }
                )
        if spec.content_box_in:
            x, y, w, h = spec.content_box_in
            if x + w > width_in + 1e-6 or y + h > height_in + 1e-6:
                errors.append(
                    {
                        "kind": kind,
                        "msg": f"content_box_in {spec.content_box_in} exceeds the slide",
                        "slide_in": [round(width_in, 3), round(height_in, 3)],
                    }
                )
    if errors:
        raise ValidationFailed("Template map does not match the template", errors)
