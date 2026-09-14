"""PPTX report builder (python-pptx): one slide per SlideSpec, laid out by a template map, rendered from the snapshot.

Rules (spec section 6):
* Every number comes from `render_view(snapshot, ai_mode)`; `--ai none` drops AI-derived facts and tables (their
  slides show an exclusion note) and AI run lines.
* table, findings, attention_list, provenance and kpis paginate; `when: has_rows` with no rows gives no slide.
* Charts are native, editable PowerPoint charts. Placeholders the slide does not fill are removed.
* Every slide carries the footer and classification label, a SYNTHETIC banner on synthetic data, a DRAFT stamp in
  draft mode, and speaker notes with sources, as-of dates, snapshot id and sha, SLA source and AI run lines.
* The final provenance slide lists freshness, data_as_of and period_end, the snapshot, AI runs
  (`format_provenance_line`), suppressed/acknowledged rule findings and omitted sections.
* Narrative (AI prose) slides are omitted in M2 and listed as omitted; `template_proof` renders them with dummy text.

Ticket and contract text in snapshot tables is untrusted data: it is only ever placed as plain text.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from sed.errors import ValidationFailed
from sed.reports.snapshot import RenderView, Snapshot, format_provenance_line, render_view
from sed.reports.specs import KpiSpec, ReportSpec, SlideSpec
from sed.reports.template_map import EMU_PER_INCH, LayoutSpec, LoadedTemplateMap, open_template, resolve_layout

PAGINATED_KINDS = frozenset({"table", "findings", "attention_list", "provenance", "kpis"})
TABLE_KINDS = frozenset({"table", "findings", "attention_list"})
CHART_KINDS = frozenset({"line_chart", "bar_chart", "stacked_bar"})
NUMERIC_UNITS = frozenset({"count", "number", "pct", "pp", "ratio", "hours", "eur"})
SYNTHETIC_BANNER = "SYNTHETIC DATA"
SYNTHETIC_NOTICE = "SYNTHETIC DATA - generated test data, not for distribution"
DRAFT_STAMP = "DRAFT"
EXCLUDED_TEXT = "AI-derived content excluded (--ai none)."
NO_ROWS_TEXT = "No rows for this period."
AI_ASSISTED_TEXT = "AI-assisted: see the provenance slide for runs and sample accuracy."
NARRATIVE_OMITTED = "AI narrative is not rendered in M2"
SHAPE_PREFIX = "sed-"

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_FACT_TOKEN = re.compile(r"\{\{\s*f:([A-Za-z0-9_.:\-]+)\s*\}\}")
_NAMED_TOKEN = re.compile(r"\{(report_title|period|as_of|data_as_of|period_end|vendor|data_class)\}")
_NUMBER_FORMATS = {
    "count": "#,##0",
    "number": "#,##0.0",
    "pct": '0.0"%"',
    "pp": '+0.0" pp";-0.0" pp"',
    "ratio": "0%",
    "hours": '0.0" h"',
    "eur": "€#,##0",
}
_SEVERITY_COLORS = {"critical": "B02A37", "high": "B02A37", "medium": "C55A11", "low": "595959"}
_MUTED = "595959"
_TILE_FILL = "F2F2F2"


# -- formatting ----------------------------------------------------------------------------------------------------


def clean_text(value: Any, *, single_line: bool = False) -> str:
    """Plain text safe for OOXML (control characters removed)."""
    text = _ILLEGAL_XML.sub("", "" if value is None else str(value))
    if single_line:
        text = " ".join(text.split())
    return text


def format_value(value: Any, unit: str, *, empty: str = "") -> str:
    """Format a fact or cell value by unit (count, pct, pp, ratio, hours, eur, date, datetime, number, text)."""
    if value is None or value == "":
        return empty
    if isinstance(value, bool):
        value = int(value)
    if unit in NUMERIC_UNITS:
        try:
            v = float(value)
        except (TypeError, ValueError):
            return clean_text(value, single_line=True)
        if math.isnan(v):
            return empty
        if unit == "count":
            return f"{round(v):,}"
        if unit == "pct":
            return f"{v:.1f}%"
        if unit == "pp":
            return f"{v:+.1f} pp"
        if unit == "ratio":
            return f"{100 * v:.0f}%"
        if unit == "hours":
            return f"{v:,.1f} h"
        if unit == "eur":
            return f"€{v:,.0f}"
        return f"{v:,.1f}"
    text = clean_text(value, single_line=True)
    if unit == "datetime":
        return text[:16].replace("T", " ")
    if unit == "date":
        return text[:10]
    return text


def _fact_text(fact: dict[str, Any] | None) -> str:
    if not fact:
        return "n/a"
    return format_value(fact.get("value"), fact.get("unit", "text"), empty="n/a")


def _truncate(text: str, limit: int) -> str:
    if limit <= 1 or len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, str):
        try:
            return float(value) if value not in (None, "") else None
        except ValueError:
            return None
    try:
        v = float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(v) else v


# -- planning ------------------------------------------------------------------------------------------------------


@dataclass
class PlannedSlide:
    """One rendered slide: the spec it comes from, its page and what it shows."""

    spec: SlideSpec
    spec_index: int
    page: int = 1
    pages: int = 1
    state: str = "content"  # content | excluded | empty
    rows: list[dict[str, Any]] = field(default_factory=list)
    kpis: list[KpiSpec] = field(default_factory=list)
    lines: list[tuple[str, bool]] = field(default_factory=list)  # provenance: (text, is_heading)
    bullets: list[str] = field(default_factory=list)  # narrative (template proof only)


@dataclass
class DeckPlan:
    slides: list[PlannedSlide]
    omitted: list[str]
    view: RenderView


def effective_slides(spec: ReportSpec, snapshot: Snapshot | None = None) -> list[SlideSpec]:
    """The spec's slides, or a default deck (title, KPIs, one chart/table per sheet) when it declares none.

    A provenance slide is appended when the spec has none: every deck ends with provenance.
    """
    slides = list(spec.slides)
    if not slides:
        slides.append(SlideSpec(kind="title"))
        facts = snapshot.facts if snapshot is not None else None
        kpis = []
        for k in spec.kpis:
            if facts is not None and k.fact not in facts:
                continue
            update = {}
            if facts is not None:
                update = {
                    name: None for name in ("compare", "delta") if getattr(k, name) and getattr(k, name) not in facts
                }
            kpis.append(k.model_copy(update=update))
        if kpis:
            slides.append(SlideSpec(kind="kpis", title="Key indicators", kpis=kpis))
        for sheet in spec.sheets:
            tbl = snapshot.tables.get(sheet.table) if snapshot is not None else None
            if snapshot is not None and tbl is None:
                continue
            keys = {c["key"] for c in tbl["columns"]} if tbl else None
            chart = sheet.chart
            if chart and (keys is None or {chart.categories, *chart.series} <= keys):
                kind = "line_chart" if chart.type == "line" else "bar_chart"
                slides.append(SlideSpec(kind=kind, title=chart.title or sheet.sheet, table=sheet.table, chart=chart))
            slides.append(SlideSpec(kind="table", title=sheet.sheet, table=sheet.table))
    if not any(s.kind == "provenance" for s in slides):
        slides.append(SlideSpec(kind="provenance", title="Provenance"))
    return slides


def check_spec_refs(snapshot: Snapshot, spec: ReportSpec) -> None:
    """Every fact, table, column and chart series a slide references must exist in the snapshot (else exit 2)."""
    errors: list[dict[str, str]] = []
    for i, s in enumerate(effective_slides(spec, snapshot)):
        loc = f"slides[{i}]"
        for k in s.kpis or []:
            for key in (k.fact, k.compare, k.delta):
                if key and key not in snapshot.facts:
                    errors.append({"loc": loc, "msg": f"fact '{key}' is not in the snapshot"})
        if not s.table or s.kind in {"title", "section", "narrative", "provenance", "kpis"}:
            continue
        tbl = snapshot.tables.get(s.table)
        if tbl is None:
            errors.append({"loc": loc, "msg": f"table '{s.table}' is not in the snapshot"})
            continue
        keys = {c["key"] for c in tbl["columns"]}
        wanted = list(s.columns or [])
        if s.chart:
            wanted += [s.chart.categories, *s.chart.series]
        for key in wanted:
            if key not in keys:
                errors.append({"loc": loc, "msg": f"column '{key}' is not in table '{s.table}'"})
    if errors:
        raise ValidationFailed(f"Report spec '{spec.report}' references content missing from the snapshot", errors)


def _pages(n: int, per_page: int) -> int:
    return max(1, math.ceil(n / per_page))


def _slide_title(s: SlideSpec, snapshot: Snapshot, spec: ReportSpec) -> str:
    if s.title:
        return s.title
    if s.kind == "title":
        return spec.title
    if s.table and s.table in snapshot.tables:
        return str(snapshot.tables[s.table].get("title") or s.table)
    return {"kpis": "Key indicators", "provenance": "Provenance", "narrative": s.ai_section_key or ""}.get(s.kind, "")


def plan_deck(
    snapshot: Snapshot,
    spec: ReportSpec,
    tmap: LoadedTemplateMap,
    *,
    ai_mode: str,
    generated_at: str,
    narratives: dict[str, list[str]] | None = None,
) -> DeckPlan:
    """Decide every slide and page before anything is drawn (validates spec references first)."""
    check_spec_refs(snapshot, spec)
    view = render_view(snapshot, ai_mode)
    slides = effective_slides(spec, snapshot)
    planned: list[PlannedSlide | int] = []  # int placeholders mark provenance positions
    omitted: list[str] = []
    for idx, s in enumerate(slides):
        layout = tmap.map.layouts[s.kind]
        per_page = s.max_rows or layout.max_rows
        title = _slide_title(s, snapshot, spec)
        if s.kind == "provenance":
            planned.append(idx)
            continue
        if s.kind == "narrative":
            if narratives is not None and s.ai_section_key in narratives:
                planned.append(PlannedSlide(s, idx, bullets=list(narratives[s.ai_section_key])))
            else:
                omitted.append(f"{title or s.ai_section_key}: {NARRATIVE_OMITTED} (ai_section_key {s.ai_section_key})")
            continue
        if s.kind == "kpis":
            visible = [k for k in s.kpis or [] if k.fact in view.facts]
            if not visible:
                planned.append(PlannedSlide(s, idx, state="excluded"))
                continue
            pages = _pages(len(visible), per_page)
            for p in range(pages):
                chunk = visible[p * per_page : (p + 1) * per_page]
                planned.append(PlannedSlide(s, idx, page=p + 1, pages=pages, kpis=chunk))
            continue
        if s.kind in TABLE_KINDS or s.kind in CHART_KINDS:
            source_rows = snapshot.tables[s.table]["rows"] if s.table in snapshot.tables else []
            if s.table in view.excluded_tables:
                if s.when == "has_rows" and not source_rows:
                    omitted.append(f"{title}: no rows this period")
                    continue
                omitted.append(f"{title}: {EXCLUDED_TEXT}")
                planned.append(PlannedSlide(s, idx, state="excluded"))
                continue
            rows = list(view.tables[s.table]["rows"])
            if not rows:
                if s.when == "has_rows":
                    omitted.append(f"{title}: no rows this period")
                    continue
                planned.append(PlannedSlide(s, idx, state="empty"))
                continue
            if s.kind in CHART_KINDS:
                planned.append(PlannedSlide(s, idx, rows=rows[: s.max_rows] if s.max_rows else rows))
                continue
            pages = _pages(len(rows), per_page)
            for p in range(pages):
                chunk = rows[p * per_page : (p + 1) * per_page]
                planned.append(PlannedSlide(s, idx, page=p + 1, pages=pages, rows=chunk))
            continue
        planned.append(PlannedSlide(s, idx))

    lines = provenance_lines(snapshot, spec, view, tmap, ai_mode=ai_mode, generated_at=generated_at, omitted=omitted)
    out: list[PlannedSlide] = []
    for item in planned:
        if not isinstance(item, int):
            out.append(item)
            continue
        s = slides[item]
        layout = tmap.map.layouts["provenance"]
        chunks = _paginate_lines(lines, s.max_rows or layout.max_rows, layout, tmap)
        for p, chunk in enumerate(chunks):
            out.append(PlannedSlide(s, item, page=p + 1, pages=len(chunks), lines=chunk))
    return DeckPlan(out, omitted, view)


def provenance_lines(
    snapshot: Snapshot,
    spec: ReportSpec,
    view: RenderView,
    tmap: LoadedTemplateMap,
    *,
    ai_mode: str,
    generated_at: str,
    omitted: list[str],
) -> list[tuple[str, bool]]:
    """Lines of the provenance slide as (text, is_heading)."""
    lines: list[tuple[str, bool]] = [("Snapshot", True)]
    lines.append((f"Data class: {snapshot.data_class.upper()}", False))
    period = f"{spec.title} | period {snapshot.period}"
    if snapshot.vendor_id:
        period += f" | vendor {snapshot.vendor_id}"
    lines.append((f"{period} | as of {snapshot.as_of}", False))
    lines.append(
        (
            f"Data as of: {snapshot.data_as_of or 'n/a'} | period end (exclusive): {snapshot.period_end or 'n/a'}"
            f" | reporting timezone {snapshot.reporting_tz}",
            False,
        )
    )
    lines.append((f"Snapshot id: {snapshot.snapshot_id}", False))
    lines.append((f"Snapshot sha256: {snapshot.sha256}", False))
    lines.append((f"SLA source: {snapshot.sla_source or 'n/a'}", False))
    lines.append(
        (
            f"Generated at (UTC): {generated_at} | code version {snapshot.git_commit or 'n/a'}"
            f" | template map {tmap.map.name} ({tmap.sha256[:12]})",
            False,
        )
    )
    lines.append(("AI runs", True))
    mode = f"AI content mode: {ai_mode}"
    if view.note:
        mode += f" ({view.note})"
    lines.append((mode, False))
    if ai_mode == "none":
        lines.append(("AI runs: excluded (--ai none)" if snapshot.ai_runs else "AI runs: none", False))
    elif not view.ai_runs:
        lines.append(("AI runs: none", False))
    lines += [(format_provenance_line(run), False) for run in view.ai_runs]
    lines.append(("Suppressed/acknowledged rule findings", True))
    if not snapshot.suppressed_findings:
        lines.append(("none", False))
    for sf in snapshot.suppressed_findings:
        until = f", suppressed until {sf['suppress_until']}" if sf.get("suppress_until") else ""
        lines.append(
            (clean_text(f"{sf.get('kind')}: {sf.get('title')} ({sf.get('status')}{until})", single_line=True), False)
        )
    lines.append(("Omitted sections", True))
    lines += [(clean_text(item, single_line=True), False) for item in omitted] or [("none", False)]
    lines.append(("Data freshness", True))
    if not snapshot.freshness:
        lines.append(("no imports recorded", False))
    for fr in snapshot.freshness:
        parts = [f"last import {fr.get('last_import') or 'n/a'}"]
        if fr.get("files") is not None:
            parts.append(f"{fr['files']} file(s)")
        if fr.get("latest_as_of"):
            parts.append(f"latest as-of {fr['latest_as_of']}")
        lines.append((clean_text(f"{fr.get('mapping_name')}: {', '.join(parts)}", single_line=True), False))
    return lines


def _text_metrics(box_in: list[float], pt: float) -> tuple[int, int]:
    """(characters per visual line, visual lines that fit) for a text box at a font size."""
    _, _, w, h = box_in
    per_line = max(10, int((w - 0.2) * 72 / (pt * 0.52)))
    capacity = max(1, int((h - 0.1) * 72 / (pt * 1.2)))
    return per_line, capacity


def _provenance_pt(tmap: LoadedTemplateMap) -> float:
    return max(8.0, tmap.map.fonts.body_pt - 4)


def _paginate_lines(
    lines: list[tuple[str, bool]], per_page: int, layout: LayoutSpec, tmap: LoadedTemplateMap
) -> list[list[tuple[str, bool]]]:
    """Split lines into pages by line count (`max_rows`) and by the visual lines the content box holds.

    A heading never ends a page, and a section that continues on the next page repeats its heading.
    """
    per_line, capacity = _text_metrics(layout.content_box_in or [0.5, 1.5, 9.0, 5.5], _provenance_pt(tmap))
    pages: list[list[tuple[str, bool]]] = []
    page: list[tuple[str, bool]] = []
    used = 0
    current_heading = ""
    for raw, heading in lines:
        text = _truncate(raw, per_line * 3)
        cost = max(1, math.ceil(len(text) / per_line))
        if heading:
            current_heading = text
        if page and (len(page) >= per_page or used + cost > capacity or (heading and len(page) >= per_page - 1)):
            pages.append(page)
            page, used = [], 0
            if not heading and current_heading and per_page >= 3:
                page.append((f"{current_heading} (continued)", True))
                used += 1
        page.append((text, heading))
        used += cost
    if page or not pages:
        pages.append(page)
    return pages


# -- rendering -----------------------------------------------------------------------------------------------------


def _emu(inches: float) -> int:
    return round(inches * EMU_PER_INCH)


class _Deck:
    def __init__(self, snapshot: Snapshot, spec: ReportSpec, tmap: LoadedTemplateMap, ai_mode: str) -> None:
        from sed.modules import metric_definitions

        self.snapshot = snapshot
        self.spec = spec
        self.loaded = tmap
        self.tmap = tmap.map
        self.ai_mode = ai_mode
        self.prs = open_template(tmap)
        self.width = self.prs.slide_width
        self.height = self.prs.slide_height
        self.definitions = metric_definitions()
        self.title_pt = _master_title_pt(self.prs)
        self.manifest: list[dict[str, Any]] = []

    # -- text helpers --------------------------------------------------------------------------------------------

    def resolve(self, text: str, view: RenderView) -> str:
        values = {
            "report_title": self.spec.title,
            "period": self.snapshot.period,
            "as_of": self.snapshot.as_of,
            "data_as_of": self.snapshot.data_as_of or "n/a",
            "period_end": self.snapshot.period_end or "n/a",
            "vendor": self.snapshot.vendor_id or "",
            "data_class": self.snapshot.data_class.upper(),
        }
        text = _FACT_TOKEN.sub(lambda m: _fact_text(view.facts.get(m.group(1))), text or "")
        return clean_text(_NAMED_TOKEN.sub(lambda m: str(values[m.group(1)]), text))

    def _font(
        self,
        run: Any,
        *,
        size: float | None = None,
        bold: bool | None = None,
        color: str | None = None,
        title: bool = False,
    ) -> None:
        from pptx.dml.color import RGBColor
        from pptx.util import Pt

        run.font.name = self.tmap.fonts.title if title else self.tmap.fonts.body
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def _write(self, text_frame: Any, paragraphs: list[dict[str, Any]], *, title: bool = False) -> None:
        """Replace a text frame's content; each paragraph dict has text and optional size/bold/color/level/runs."""
        from pptx.util import Pt

        text_frame.clear()
        text_frame.word_wrap = True
        for i, para in enumerate(paragraphs):
            p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
            if para.get("level"):
                p.level = para["level"]
            if para.get("align") is not None:
                p.alignment = para["align"]
            if para.get("space_after") is not None:
                p.space_after = Pt(para["space_after"])
            runs = para.get("runs") or [para]
            for spec in runs:
                run = p.add_run()
                run.text = clean_text(spec.get("text", ""))
                self._font(
                    run,
                    size=spec.get("size", para.get("size")),
                    bold=spec.get("bold", para.get("bold")),
                    color=spec.get("color", para.get("color")),
                    title=title,
                )

    def _textbox(self, slide: Any, box: list[float], paragraphs: list[dict[str, Any]], name: str) -> Any:
        shape = slide.shapes.add_textbox(_emu(box[0]), _emu(box[1]), _emu(box[2]), _emu(box[3]))
        shape.name = SHAPE_PREFIX + name
        self._write(shape.text_frame, paragraphs)
        return shape

    # -- slide scaffolding ---------------------------------------------------------------------------------------

    def new_slide(self, kind: str, title: str) -> tuple[Any, set[int], LayoutSpec]:
        lspec = self.tmap.layouts[kind]
        layout, fallback = resolve_layout(self.prs, lspec)
        slide = self.prs.slides.add_slide(layout)
        used: set[int] = set()
        self.manifest.append({"kind": kind, "title": title, "layout": layout.name, "fallback_used": fallback})
        if title:
            idx = lspec.placeholders.get("title")
            ph = _placeholder(slide, idx)
            if ph is not None:
                size = _fitting_size(title, ph.width, ph.height, self.title_pt, max_lines=2)
                self._write(ph.text_frame, [{"text": title, "size": size}], title=True)
                used.add(idx)
            else:
                box = self.content_box(kind)
                top = 0.35
                height = max(0.4, min(1.0, box[1] - top - 0.05)) if box[1] > top + 0.45 else 0.4
                shape = self._textbox(
                    slide,
                    [box[0], top, box[2], height],
                    [{"text": title, "size": min(28, self.tmap.fonts.body_pt + 10), "bold": True}],
                    "title",
                )
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.name = self.tmap.fonts.title
        return slide, used, lspec

    def content_box(self, kind: str) -> list[float]:
        box = self.tmap.layouts[kind].content_box_in
        if box:
            return list(box)
        w = self.width / EMU_PER_INCH
        h = self.height / EMU_PER_INCH
        return [0.5, 1.5, w - 1.0, h - 2.0]

    def finish(self, slide: Any, used: set[int], notes: str, number: int) -> None:
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx not in used:
                ph._element.getparent().remove(ph._element)
        w_in = self.width / EMU_PER_INCH
        h_in = self.height / EMU_PER_INCH
        footer_pt = max(7.0, self.tmap.fonts.body_pt - 5)
        footer = self.resolve(self.tmap.footer_text, RenderView({}, {}, [], [], [], None))
        left_w = max(1.0, w_in * 0.62)
        self._textbox(
            slide, [0.3, h_in - 0.42, left_w, 0.3], [{"text": footer, "size": footer_pt, "color": _MUTED}], "footer"
        )
        right = f"{self.tmap.classification_label} | {number}" if self.tmap.classification_label else str(number)
        right_w = max(0.8, w_in - 0.6 - left_w - 0.1)
        self._textbox(
            slide,
            [w_in - 0.3 - right_w, h_in - 0.42, right_w, 0.3],
            [{"text": right, "size": footer_pt, "color": _MUTED, "align": PP_ALIGN.RIGHT}],
            "classification",
        )
        stamps = []
        if self.snapshot.data_class == "synthetic":
            stamps.append(("synthetic-banner", SYNTHETIC_BANNER, [w_in - 2.05, 0.05, 2.0, 0.28], "B02A37", "FFFFFF"))
        if self.ai_mode == "draft":
            stamps.append(("draft-stamp", DRAFT_STAMP, [0.05, 0.05, 1.1, 0.28], "FFFFFF", "B02A37"))
        for name, text, box, fill, color in stamps:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(_emu(v) for v in box))
            shape.name = SHAPE_PREFIX + name
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill)
            shape.line.color.rgb = _rgb("B02A37")
            shape.shadow.inherit = False
            tf = shape.text_frame
            tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._write(tf, [{"text": text, "size": 10, "bold": True, "color": color, "align": PP_ALIGN.CENTER}])
        slide.notes_slide.notes_text_frame.text = clean_text(notes)

    # -- notes ---------------------------------------------------------------------------------------------------

    def notes(self, planned: PlannedSlide, view: RenderView) -> str:
        s = planned.spec
        snap = self.snapshot
        out: list[str] = []
        if s.notes:
            out.append(self.resolve(s.notes, view))
        if s.table:
            tbl = snap.tables.get(s.table) or {}
            out.append(f"Source: snapshot table '{s.table}' ({tbl.get('title', '')})")
        if planned.kpis:
            out.append("Source: snapshot facts " + ", ".join(k.fact for k in planned.kpis))
            defs = []
            for k in planned.kpis:
                for key in (k.fact, k.compare, k.delta):
                    f = view.facts.get(key or "")
                    if f and f.get("definition") and f["definition"] in self.definitions:
                        entry = f"- {f['definition']}: {self.definitions[f['definition']][1]}"
                        if entry not in defs:
                            defs.append(entry)
            if defs:
                out.append("Definitions:\n" + "\n".join(defs))
        mappings = [str(fr.get("mapping_name")) for fr in snap.freshness if fr.get("files") is not None]
        if mappings:
            out.append("Imported sources: " + ", ".join(mappings))
        out.append(
            f"As of {snap.as_of} | data as of {snap.data_as_of or 'n/a'} | period end (exclusive) "
            f"{snap.period_end or 'n/a'}"
        )
        out.append(f"Snapshot {snap.snapshot_id} sha256 {snap.sha256}")
        out.append(f"SLA source: {snap.sla_source or 'n/a'}")
        if s.table and s.table in snap.ai_derived_tables and self.ai_mode != "none":
            out.append(AI_ASSISTED_TEXT)
        if self.ai_mode == "none":
            out.append("AI runs: excluded (--ai none)" if snap.ai_runs else "AI runs: none")
        else:
            out.append("AI runs:" if view.ai_runs else "AI runs: none")
            out += [format_provenance_line(run) for run in view.ai_runs]
        out.append(f"AI content mode: {self.ai_mode}; data class {snap.data_class.upper()}")
        return "\n".join(line for line in out if line)

    # -- kinds ---------------------------------------------------------------------------------------------------

    def render(self, planned: PlannedSlide, view: RenderView, number: int) -> None:
        s = planned.spec
        title = self.resolve(_slide_title(s, self.snapshot, self.spec), view)
        if planned.pages > 1:
            title = f"{title} ({planned.page}/{planned.pages})"
        slide, used, lspec = self.new_slide(s.kind, title)
        box = self.content_box(s.kind)
        if planned.state == "excluded":
            self._textbox(slide, box, [{"text": EXCLUDED_TEXT, "size": self.tmap.fonts.body_pt}], "excluded-note")
        elif planned.state == "empty":
            self._textbox(slide, box, [{"text": NO_ROWS_TEXT, "size": self.tmap.fonts.body_pt}], "no-rows")
        elif s.kind == "title":
            self._title_body(slide, used, lspec, s, view)
        elif s.kind == "section":
            self._section_body(slide, used, lspec, s, view)
        elif s.kind == "narrative":
            self._narrative(slide, used, lspec, planned)
        elif s.kind == "kpis":
            self._kpis(slide, box, planned.kpis, view)
        elif s.kind in CHART_KINDS:
            if s.table in self.snapshot.ai_derived_tables:
                box = self._caption(slide, box)
            self._chart(slide, box, s, view.tables[s.table], planned.rows)
        elif s.kind in TABLE_KINDS:
            if s.table in self.snapshot.ai_derived_tables:
                box = self._caption(slide, box)
            tbl = view.tables[s.table]
            columns = s.columns or [c["key"] for c in tbl["columns"]]
            if s.kind == "findings":
                self._findings(slide, box, tbl, columns, planned.rows)
            else:
                self._table(slide, box, tbl, columns, planned.rows, emphasis=s.kind == "attention_list")
        elif s.kind == "provenance":
            self._provenance(slide, box, planned.lines)
        self.finish(slide, used, self.notes(planned, view), number)

    def _caption(self, slide: Any, box: list[float]) -> list[float]:
        size = max(8.0, self.tmap.fonts.body_pt - 5)
        self._textbox(
            slide, [box[0], box[1] + box[3] - 0.3, box[2], 0.3], [{"text": AI_ASSISTED_TEXT, "size": size}], "ai-note"
        )
        return [box[0], box[1], box[2], max(0.5, box[3] - 0.35)]

    def _body_target(self, slide: Any, used: set[int], lspec: LayoutSpec, role: str, kind: str) -> Any:
        idx = lspec.placeholders.get(role)
        ph = _placeholder(slide, idx)
        if ph is not None:
            used.add(idx)
            return ph.text_frame
        w_in = self.width / EMU_PER_INCH
        h_in = self.height / EMU_PER_INCH
        box = self.tmap.layouts[kind].content_box_in or [0.75, h_in * 0.55, w_in - 1.5, h_in * 0.3]
        shape = slide.shapes.add_textbox(*(_emu(v) for v in box))
        shape.name = SHAPE_PREFIX + role
        return shape.text_frame

    def _title_body(self, slide: Any, used: set[int], lspec: LayoutSpec, s: SlideSpec, view: RenderView) -> None:
        snap = self.snapshot
        lines = [self.resolve(s.subtitle, view) if s.subtitle else f"Period {snap.period}"]
        if snap.vendor_id:
            lines.append(f"Vendor {snap.vendor_id}")
        lines.append(f"As of {snap.as_of} | data as of {snap.data_as_of or 'n/a'}")
        if self.spec.audience:
            lines.append(self.spec.audience)
        if snap.data_class == "synthetic":
            lines.append(SYNTHETIC_NOTICE)
        if self.ai_mode == "draft":
            lines.append("DRAFT - includes unapproved AI content")
        size = max(10.0, self.tmap.fonts.body_pt + 2)
        paragraphs = [{"text": line, "size": size} for line in lines]
        if snap.data_class == "synthetic":
            paragraphs[lines.index(SYNTHETIC_NOTICE)].update({"bold": True, "color": "B02A37"})
        self._write(self._body_target(slide, used, lspec, "subtitle", "title"), paragraphs)

    def _section_body(self, slide: Any, used: set[int], lspec: LayoutSpec, s: SlideSpec, view: RenderView) -> None:
        if not s.subtitle:
            return
        text = self.resolve(s.subtitle, view)
        self._write(self._body_target(slide, used, lspec, "body", "section"), [{"text": text}])

    def _narrative(self, slide: Any, used: set[int], lspec: LayoutSpec, planned: PlannedSlide) -> None:
        bullets = planned.bullets[: lspec.max_bullets]
        if not bullets:
            return
        paragraphs = [{"text": b, "size": self.tmap.fonts.body_pt + 2} for b in bullets]
        self._write(self._body_target(slide, used, lspec, "body", "narrative"), paragraphs)

    def _kpis(self, slide: Any, box: list[float], kpis: list[KpiSpec], view: RenderView) -> None:
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR

        n = len(kpis)
        cols = {1: 1, 2: 2, 3: 3, 4: 2, 5: 3, 6: 3, 9: 3}.get(n, 4)
        rows = math.ceil(n / cols)
        gap = 0.15
        x0, y0, bw, bh = box
        tile_w = (bw - gap * (cols - 1)) / cols
        tile_h = min((bh - gap * (rows - 1)) / rows, 2.2)
        body = self.tmap.fonts.body_pt
        value_pt = max(14.0, min(30.0, tile_h * 72 * 0.28, tile_w * 72 / 6))
        small = max(8.0, body - 5)
        for i, k in enumerate(kpis):
            r, c = divmod(i, cols)
            x = x0 + c * (tile_w + gap)
            y = y0 + r * (tile_h + gap)
            fact = view.facts[k.fact]
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), _emu(y), _emu(tile_w), _emu(tile_h))
            shape.name = f"{SHAPE_PREFIX}kpi-{k.fact}"
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(_TILE_FILL)
            shape.line.color.rgb = _rgb(self.tmap.series_colors[0][1:])
            shape.shadow.inherit = False
            tf = shape.text_frame
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = _emu(0.08)
            tf.margin_top = tf.margin_bottom = _emu(0.05)
            label_pt = max(8.0, body - 3)
            paragraphs = [
                {
                    "text": _truncate(clean_text(fact.get("label"), single_line=True), 60),
                    "size": label_pt,
                    "color": "404040",
                },
                {"text": _fact_text(fact), "size": value_pt, "bold": True, "color": self.tmap.series_colors[0][1:]},
            ]
            for key in (k.compare, k.delta):
                other = view.facts.get(key or "")
                if other:
                    label = _truncate(clean_text(other.get("label"), single_line=True), 40)
                    paragraphs.append({"text": f"{label}: {_fact_text(other)}", "size": small, "color": _MUTED})
            self._write(tf, paragraphs)

    def _chart(self, slide: Any, box: list[float], s: SlideSpec, tbl: dict[str, Any], rows: list[dict]) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Pt

        chart_spec = s.chart
        assert chart_spec is not None  # guaranteed by SlideSpec validation
        columns = {c["key"]: c for c in tbl["columns"]}
        unit = columns[chart_spec.series[0]].get("format", "number")
        number_format = _NUMBER_FORMATS.get(unit, "General")
        data = CategoryChartData(number_format=number_format)
        cat_unit = columns[chart_spec.categories].get("format", "text")
        data.categories = [format_value(r.get(chart_spec.categories), cat_unit, empty="(blank)") for r in rows]
        for key in chart_spec.series:
            data.add_series(clean_text(columns[key]["label"]), [_number(r.get(key)) for r in rows])
        horizontal = chart_spec.type == "bar"
        if s.kind == "line_chart" or (s.kind == "bar_chart" and chart_spec.type == "line"):
            chart_type = XL_CHART_TYPE.LINE_MARKERS
        elif s.kind == "stacked_bar":
            chart_type = XL_CHART_TYPE.BAR_STACKED if horizontal else XL_CHART_TYPE.COLUMN_STACKED
        else:
            chart_type = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        frame = slide.shapes.add_chart(chart_type, *(_emu(v) for v in box), data)
        frame.name = f"{SHAPE_PREFIX}chart-{s.table}"
        chart = frame.chart
        body = self.tmap.fonts.body_pt
        chart.font.name = self.tmap.fonts.body
        chart.font.size = Pt(max(8.0, body - 4))
        chart.has_legend = len(chart_spec.series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        if chart_spec.title:
            chart.has_title = True
            chart.chart_title.text_frame.text = clean_text(chart_spec.title)
            for p in chart.chart_title.text_frame.paragraphs:
                for run in p.runs:
                    self._font(run, size=max(9.0, body - 2), bold=True)
        else:
            chart.has_title = False
        plot = chart.plots[0]
        line = chart_type == XL_CHART_TYPE.LINE_MARKERS
        for i, series in enumerate(plot.series):
            color = _rgb(self.tmap.series_colors[i % len(self.tmap.series_colors)][1:])
            if line:
                series.format.line.color.rgb = color
                series.format.line.width = Pt(2.25)
                series.smooth = False
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
        if horizontal and not line:
            chart.category_axis.reverse_order = True
        if s.kind == "bar_chart" and not line and len(rows) <= 12 and len(chart_spec.series) <= 2:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format = number_format
            labels.number_format_is_linked = False
            labels.font.size = Pt(max(8.0, body - 5))

    def _table(
        self, slide: Any, box: list[float], tbl: dict[str, Any], columns: list[str], rows: list[dict], *, emphasis: bool
    ) -> None:
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

        col_defs = {c["key"]: c for c in tbl["columns"]}
        cols = [col_defs[k] for k in columns]
        pt = self.tmap.fonts.table_pt
        x, y, w, h = box
        n_rows = len(rows) + 1
        line_h = pt * 1.2 / 72
        margin_v = 0.06
        row_avail = h / n_rows
        max_lines = max(1, min(3, int((row_avail - margin_v) / line_h)))
        row_h = min(row_avail, max_lines * line_h + margin_v + 0.02)
        all_rows = tbl["rows"]
        weights = []
        for c in cols:
            longest = max([len(format_value(r.get(c["key"]), c.get("format", "text"))) for r in all_rows] or [0])
            weights.append(min(45, max(4, len(str(c["label"])), longest)))
        total = sum(weights)
        widths = [int(_emu(w) * wt / total) for wt in weights]
        widths[-1] = _emu(w) - sum(widths[:-1])
        frame = slide.shapes.add_table(n_rows, len(cols), _emu(x), _emu(y), _emu(w), _emu(row_h * n_rows))
        frame.name = f"{SHAPE_PREFIX}table"
        table = frame.table
        for i, width in enumerate(widths):
            table.columns[i].width = width
        header_color = self.tmap.series_colors[0][1:]
        for ci, c in enumerate(cols):
            per_line = max(3, int(((widths[ci] / EMU_PER_INCH) - 0.1) * 72 / (pt * 0.52)))
            numeric = c.get("format") in NUMERIC_UNITS
            header = table.cell(0, ci)
            header.fill.solid()
            header.fill.fore_color.rgb = _rgb(header_color)
            label = _truncate(clean_text(c["label"], single_line=True), per_line * max_lines)
            self._cell(header, label, pt, bold=True, color="FFFFFF", align=PP_ALIGN.RIGHT if numeric else None)
            for ri, row in enumerate(rows, start=1):
                text = _truncate(format_value(row.get(c["key"]), c.get("format", "text")), per_line * max_lines)
                cell = table.cell(ri, ci)
                self._cell(cell, text, pt, bold=emphasis and ci == 0, align=PP_ALIGN.RIGHT if numeric else None)
        for r in range(n_rows):
            for ci in range(len(cols)):
                cell = table.cell(r, ci)
                cell.margin_left = cell.margin_right = _emu(0.05)
                cell.margin_top = cell.margin_bottom = _emu(margin_v / 2)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    def _cell(
        self, cell: Any, text: str, pt: float, *, bold: bool = False, color: str | None = None, align: Any = None
    ) -> None:
        self._write(cell.text_frame, [{"text": text, "size": pt, "bold": bold, "color": color, "align": align}])

    def _findings(
        self, slide: Any, box: list[float], tbl: dict[str, Any], columns: list[str], rows: list[dict]
    ) -> None:
        col_defs = {c["key"]: c for c in tbl["columns"]}
        tag_col = "severity" if "severity" in columns else None
        head_col = "title" if "title" in columns else next((c for c in columns if c != tag_col), columns[0])
        detail_cols = [c for c in columns if c not in {tag_col, head_col}]
        body = self.tmap.fonts.body_pt
        detail_pt = max(8.0, body - 4)
        per_row_pt = box[3] * 72 / max(1, len(rows))
        head_pt = max(9.0, min(body, (per_row_pt - 6 - detail_pt * 1.2) / (2 * 1.2)))
        per_line = max(10, int((box[2] - 0.2) * 72 / (head_pt * 0.52)))
        paragraphs: list[dict[str, Any]] = []
        for row in rows:
            runs = []
            if tag_col:
                severity = format_value(row.get(tag_col), "text")
                color = _SEVERITY_COLORS.get(severity.lower(), _MUTED)
                runs.append({"text": f"[{severity.upper()}] ", "bold": True, "color": color, "size": head_pt})
            head = format_value(row.get(head_col), col_defs[head_col].get("format", "text"))
            runs.append({"text": _truncate(head, per_line * 2 - 12), "size": head_pt})
            paragraphs.append({"runs": runs, "size": head_pt, "space_after": 0})
            details = [format_value(row.get(c), col_defs[c].get("format", "text")) for c in detail_cols]
            detail = " · ".join(d for d in details if d)
            paragraphs.append(
                {"text": _truncate(detail, per_line * 2) or " ", "size": detail_pt, "color": _MUTED, "space_after": 6}
            )
        self._textbox(slide, box, paragraphs, "findings")

    def _provenance(self, slide: Any, box: list[float], lines: list[tuple[str, bool]]) -> None:
        pt = _provenance_pt(self.loaded)
        paragraphs = [
            {"text": text, "size": pt + (1 if heading else 0), "bold": heading, "color": None if heading else "262626"}
            for text, heading in lines
        ]
        self._textbox(slide, box, paragraphs, "provenance")


def _master_title_pt(prs: Any) -> float:
    """The template's title size (slide master title style), 44 pt when the template does not say."""
    from pptx.oxml.ns import qn

    styles = prs.slide_master._element.find(qn("p:txStyles"))
    level = styles.find(qn("p:titleStyle")).find(qn("a:lvl1pPr")) if styles is not None else None
    run_props = level.find(qn("a:defRPr")) if level is not None else None
    size = run_props.get("sz") if run_props is not None else None
    return int(size) / 100 if size and size.isdigit() else 44.0


def _fitting_size(
    text: str, width_emu: int | None, height_emu: int | None, default_pt: float, *, max_lines: int
) -> float | None:
    """None when the template size fits the placeholder, else the largest smaller size (>= 18 pt) that does."""
    if not width_emu or not height_emu:
        return None
    width_in = width_emu / EMU_PER_INCH - 0.2
    height_in = height_emu / EMU_PER_INCH - 0.1
    size = default_pt
    while size > 18:
        per_line = max(1, int(width_in * 72 / (size * 0.5)))
        lines = math.ceil(len(text) / per_line)
        if lines <= max_lines and lines * size * 1.2 / 72 <= height_in:
            break
        size -= 2
    size = max(size, 18.0)
    return None if size >= default_pt else size


def _placeholder(slide: Any, idx: int | None) -> Any:
    if idx is None:
        return None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def _rgb(hex_color: str) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_color.lstrip("#"))


def render_deck(
    snapshot: Snapshot,
    spec: ReportSpec,
    tmap: LoadedTemplateMap,
    out_path: Path,
    *,
    ai_mode: str,
    generated_at: str,
    narratives: dict[str, list[str]] | None = None,
) -> dict[str, Any]:
    """Render a deck and return a manifest: slides (kind, title, layout, fallback_used) and omitted sections."""
    plan = plan_deck(snapshot, spec, tmap, ai_mode=ai_mode, generated_at=generated_at, narratives=narratives)
    deck = _Deck(snapshot, spec, tmap, ai_mode)
    for number, planned in enumerate(plan.slides, start=1):
        deck.render(planned, plan.view, number)
    deck.prs.core_properties.title = clean_text(f"{spec.title} {snapshot.period}")
    deck.prs.core_properties.subject = f"snapshot {snapshot.snapshot_id}"
    deck.prs.core_properties.keywords = snapshot.data_class.upper()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(str(out_path))
    return {"path": str(out_path), "slides": deck.manifest, "omitted": plan.omitted}


def build_pptx(
    snapshot: Snapshot, spec: ReportSpec, tmap: LoadedTemplateMap, out_path: Path, *, ai_mode: str, generated_at: str
) -> Path:
    """Build the report deck for one snapshot with a loaded template map."""
    render_deck(snapshot, spec, tmap, out_path, ai_mode=ai_mode, generated_at=generated_at)
    return out_path
