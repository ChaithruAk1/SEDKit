"""Reading the trail: filters, paging, open attempts, the people in it, the check, and the recorded workbook
download."""

from __future__ import annotations

from datetime import date
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from sed.api.models import AuditPageOut
from sed.audit import store
from sed.audit.query import AuditFilters, read_page
from sed.audit.record import record, start
from sed.auth.actor import command_line_actor, developer_actor
from sed.auth.identity import Identity
from sed.paths import get_paths
from tests.fixtures.api import api_client
from tests.platform.api.conftest import assert_envelope
from tests.platform.audit.test_record import entries

UTC_TZ = ZoneInfo("UTC")
OWNER = "owner@example.com"


@pytest.fixture
def trail(data_root, monkeypatch):
    """A profile whose trail holds a signed-in download, a finished pull, an unfinished import and a formula-looking
    summary."""
    from sed import bootstrap

    monkeypatch.setenv("USERNAME", "synthetic-user")
    paths = get_paths("synthetic")
    bootstrap.init_profile(paths, write_claude_settings=False)
    owner = Identity("google", "s", emails=(OWNER,), name="Owner", email=OWNER).actor()
    record(paths, owner, "download", summary="Downloaded 3 tickets as a workbook.", detail={"rows": 3})
    with start(paths, command_line_actor(), "pull", summary="Pull from Jira") as attempt:
        attempt.done("Pull from Jira: 7 rows pulled.")
    start(paths, command_line_actor(), "import", summary="Import a.csv")  # never finished
    record(paths, developer_actor(), "download", summary="=HYPERLINK(1) looks like a formula")
    return paths


def test_the_page_lists_newest_first_with_people_actions_and_the_check(trail):
    body = read_page(trail, AuditFilters(), tz=UTC_TZ)
    assert body["total"] == 5 and [e["seq"] for e in body["items"]] == [5, 4, 3, 2, 1]
    assert body["integrity"]["intact"] and body["integrity"]["entries"] == 5
    assert {p["id"]: p["verified"] for p in body["people"]} == {OWNER: True, "windows:synthetic-user": False}
    assert {a["key"] for a in body["actions"]} >= {"download", "pull", "import", "sign_in"}
    unfinished = next(e for e in body["items"] if e["action"] == "import")
    finished = next(e for e in body["items"] if e["action"] == "pull" and e["outcome"] == "started")
    assert unfinished["open"] is True and finished["open"] is False
    AuditPageOut.model_validate(body)


def test_filters(trail):
    def seqs(**filters):
        return [e["seq"] for e in read_page(trail, AuditFilters(**filters), tz=UTC_TZ)["items"]]

    assert seqs(actor=OWNER) == [1]
    assert seqs(action="pull") == [3, 2]
    assert seqs(outcome="started") == [4, 2]
    assert seqs(q="jira") == [3, 2] and seqs(q="100%") == []  # a LIKE wildcard is matched literally
    pull = read_page(trail, AuditFilters(action="pull"), tz=UTC_TZ)["items"][0]
    assert seqs(correlation=pull["correlation_id"]) == [3, 2]
    today = date.fromisoformat(pull["at"][:10])
    assert len(seqs(since=today, until=today)) == 5 and seqs(until=date(2000, 1, 1)) == []
    page2 = read_page(trail, AuditFilters(), tz=UTC_TZ, page=2, page_size=2)
    assert [e["seq"] for e in page2["items"]] == [3, 2] and page2["total"] == 5


def test_an_empty_trail_reads_as_empty_and_intact(data_root):
    paths = get_paths("synthetic")
    body = read_page(paths, AuditFilters(), tz=UTC_TZ)
    assert body["items"] == [] and body["integrity"]["intact"] and not (paths.audit / "audit.db").exists()


def test_the_api_answers_the_page_and_refuses_bad_filters(trail):
    client = api_client(trail)
    body = AuditPageOut.model_validate(client.get("/api/audit?action=pull&page_size=1").json())
    assert body.total == 2 and len(body.items) == 1 and body.items[0].action_label == "Pull"
    assert_envelope(client.get("/api/audit?action=teleport"), 422, "validation")
    assert_envelope(client.get("/api/audit?since=2026-09-02&until=2026-09-01"), 422, "validation")
    assert_envelope(client.get("/api/audit?outcome=maybe"), 422, "validation")


def test_the_workbook_download_is_recorded_and_keeps_the_fingerprints(trail, tmp_path):
    client = api_client(trail)
    response = client.get("/api/audit-export.xlsx?action=download")
    assert response.status_code == 200 and response.headers["content-type"].startswith("application/vnd.openxmlformats")
    saved = tmp_path / "audit.xlsx"
    saved.write_bytes(response.content)
    sheet = openpyxl.load_workbook(saved).active
    header = [c.value for c in sheet[1]]
    rows = [dict(zip(header, [c.value for c in row], strict=True)) for row in sheet.iter_rows(min_row=2)]
    assert [r["Entry"] for r in rows] == [1, 5]  # oldest first
    assert (
        rows[1]["What"] == "=HYPERLINK(1) looks like a formula"
        and sheet.cell(3, header.index("What") + 1).data_type == "s"
    )
    assert rows[0]["Fingerprint"] == entries(trail)[0]["entry_hash"] and rows[0]["Proven"] == "yes"
    last = entries(trail)[-1]
    assert last["action"] == "download" and last["target_type"] == "audit_log" and last["detail"]["rows"] == 2
    assert last["detail"]["filters"] == {"action": "download"}


def test_the_workbook_is_refused_when_the_trail_cannot_record_it(trail, monkeypatch):
    def broken(path, row, **kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(store, "append", broken)
    response = api_client(trail).get("/api/audit-export.xlsx")
    assert_envelope(response, 412, "precondition")
    assert not list((trail.out / "exports").glob("audit_*.xlsx"))
