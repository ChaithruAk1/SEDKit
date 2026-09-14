"""Template map loading: schema, lookup order, .potx rejection, layout resolution, placeholders, slide size, sha."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
import yaml
from typer.testing import CliRunner

from sed.errors import ValidationFailed
from sed.paths import get_paths
from sed.reports.template_map import (
    POTX_MESSAGE,
    SLIDE_KINDS,
    default_template_path,
    load_template_map,
)
from tests.conftest import REPO

NEUTRAL = REPO / "templates" / "pptx" / "neutral.map.yaml"


def neutral_data() -> dict:
    return yaml.safe_load(NEUTRAL.read_text(encoding="utf-8"))


def write_map(path: Path, data: dict) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
    return path


def test_neutral_map_validates_and_hashes_map_plus_template():
    loaded = load_template_map("neutral")
    tmap = loaded.map
    assert tmap.name == "neutral" and tmap.template is None and tmap.slide_size == "4:3"
    assert set(tmap.layouts) == set(SLIDE_KINDS) and len(SLIDE_KINDS) == 11
    fonts = tmap.fonts
    assert (fonts.title, fonts.body, fonts.body_pt, fonts.table_pt) == ("Calibri", "Calibri", 14, 10)
    assert tmap.series_colors == ["#1F4E79", "#2E75B6", "#9DC3E6", "#C55A11", "#7F7F7F"]
    assert tmap.layouts["title"].layout_name == "Title Slide" and tmap.layouts["title"].placeholders == {
        "title": 0,
        "subtitle": 1,
    }
    assert tmap.layouts["section"].fallback_index == 2 and tmap.layouts["narrative"].placeholders == {
        "title": 0,
        "body": 1,
    }
    for kind in ("kpis", "line_chart", "bar_chart", "stacked_bar", "table", "findings", "attention_list", "provenance"):
        layout = tmap.layouts[kind]
        assert (layout.layout_name, layout.fallback_index, layout.content_box_in) == (
            "Title Only",
            5,
            [0.5, 1.5, 9.0, 5.5],
        )
    assert {k: tmap.layouts[k].max_rows for k in ("table", "findings", "attention_list", "provenance")} == {
        "table": 12,
        "findings": 8,
        "attention_list": 12,
        "provenance": 18,
    }
    expected = hashlib.sha256(NEUTRAL.read_bytes() + b"\0" + default_template_path().read_bytes()).hexdigest()
    assert loaded.sha256 == expected and loaded.template_path is None and loaded.path == NEUTRAL.resolve()


def test_default_map_comes_from_settings(tmp_path):
    paths = get_paths("synthetic", tmp_path / "profile")
    assert load_template_map(None, paths).map.name == "neutral"
    (paths.config).mkdir(parents=True)
    (paths.config / "settings.yaml").write_text("reports:\n  template_map: other\n", encoding="utf-8")
    other = neutral_data() | {"name": "other"}
    write_map(paths.config / "templates" / "other.map.yaml", other)
    assert load_template_map(None, paths).map.name == "other"


def test_missing_kind_fails(tmp_path):
    data = neutral_data()
    del data["layouts"]["provenance"]
    with pytest.raises(ValidationFailed) as exc:
        load_template_map(str(write_map(tmp_path / "broken.map.yaml", data)))
    assert exc.value.exit_code == 2 and "provenance" in json.dumps(exc.value.details)


def test_content_box_required_for_content_kinds(tmp_path):
    data = neutral_data()
    del data["layouts"]["table"]["content_box_in"]
    with pytest.raises(ValidationFailed, match="Invalid template map"):
        load_template_map(str(write_map(tmp_path / "nobox.map.yaml", data)))


def test_potx_exits_2(tmp_path):
    with pytest.raises(ValidationFailed, match=r"save as \.pptx"):
        load_template_map(str(tmp_path / "corporate.potx"))
    (tmp_path / "corporate.potx").write_bytes(b"not a real template")
    data = neutral_data() | {"template": "corporate.potx"}
    with pytest.raises(ValidationFailed) as exc:
        load_template_map(str(write_map(tmp_path / "corp.map.yaml", data)))
    assert exc.value.exit_code == 2 and exc.value.message == POTX_MESSAGE

    from sed.cli import app

    result = CliRunner().invoke(app, ["report", "template-inspect", str(tmp_path / "corporate.potx"), "--json"])
    assert result.exit_code == 2
    assert "save as .pptx" in json.loads(result.stdout.strip().splitlines()[-1])["error"]["message"]


def test_data_dir_config_templates_take_precedence(tmp_path):
    paths = get_paths("synthetic", tmp_path / "profile")
    assert load_template_map("neutral", paths).path == NEUTRAL.resolve()
    local = write_map(
        paths.config / "templates" / "neutral.map.yaml", neutral_data() | {"classification_label": "Local"}
    )
    loaded = load_template_map("neutral", paths)
    assert loaded.path == local.resolve() and loaded.map.classification_label == "Local"
    assert loaded.sha256 != load_template_map("neutral").sha256
    explicit = write_map(tmp_path / "elsewhere" / "neutral.map.yaml", neutral_data() | {"classification_label": "Path"})
    assert load_template_map(str(explicit), paths).map.classification_label == "Path"
    with pytest.raises(ValidationFailed, match="not found"):
        load_template_map("no-such-map", paths)


def test_slide_size_mismatch_exits_2(tmp_path):
    with pytest.raises(ValidationFailed, match="slide size") as exc:
        load_template_map(str(write_map(tmp_path / "wide.map.yaml", neutral_data() | {"slide_size": "16:9"})))
    assert exc.value.exit_code == 2


def test_layouts_resolve_by_name_then_fallback_index(tmp_path):
    from sed.reports.template_map import open_template, resolve_layout

    data = neutral_data()
    data["layouts"]["table"]["layout_name"] = "Layout that does not exist"
    loaded = load_template_map(str(write_map(tmp_path / "fallback.map.yaml", data)))
    layout, fallback = resolve_layout(open_template(loaded), loaded.map.layouts["table"])
    assert fallback and layout.name == "Title Only"
    layout, fallback = resolve_layout(open_template(loaded), loaded.map.layouts["kpis"])
    assert not fallback and layout.name == "Title Only"
    data["layouts"]["table"]["fallback_index"] = 99
    with pytest.raises(ValidationFailed, match="not found"):
        load_template_map(str(write_map(tmp_path / "nolayout.map.yaml", data)))


def test_declared_placeholder_must_exist_and_box_must_fit(tmp_path):
    data = neutral_data()
    data["layouts"]["narrative"]["placeholders"]["body"] = 7
    with pytest.raises(ValidationFailed) as exc:
        load_template_map(str(write_map(tmp_path / "ph.map.yaml", data)))
    assert "idx 7" in json.dumps(exc.value.details)
    data = neutral_data()
    data["layouts"]["table"]["content_box_in"] = [0.5, 1.5, 9.8, 5.5]
    with pytest.raises(ValidationFailed, match="does not match"):
        load_template_map(str(write_map(tmp_path / "box.map.yaml", data)))


def test_template_path_is_relative_to_the_map_and_hashed(tmp_path):
    from pptx import Presentation

    folder = tmp_path / "config" / "templates"
    folder.mkdir(parents=True)
    Presentation().save(str(folder / "house.pptx"))
    map_path = write_map(folder / "house.map.yaml", neutral_data() | {"name": "house", "template": "house.pptx"})
    loaded = load_template_map(str(map_path))
    assert loaded.template_path == (folder / "house.pptx").resolve()
    expected = hashlib.sha256(map_path.read_bytes() + b"\0" + (folder / "house.pptx").read_bytes()).hexdigest()
    assert loaded.sha256 == expected
    (folder / "house.pptx").unlink()
    with pytest.raises(ValidationFailed, match="Template file not found"):
        load_template_map(str(map_path))
    (folder / "house.pptx").write_bytes(b"not a zip")
    with pytest.raises(ValidationFailed, match="Cannot open"):
        load_template_map(str(map_path))
