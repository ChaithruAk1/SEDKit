"""A second template map works with no code change: renamed and reordered layouts, 16:9, other fonts and boxes.

The alternative template is generated from python-pptx's default template inside the test tmp dir (binary templates
are never committed).
"""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pytest
import yaml
from typer.testing import CliRunner

from sed.reports.build import build_report
from sed.reports.template_map import SLIDE_KINDS, load_template_map
from sed.reports.template_tools import template_proof
from tests.platform.reports.test_pptx_builder import deck_problems, open_deck, slide_text

RENAMES = {
    "Title Slide": "Corp Cover",
    "Title and Content": "Corp Text",
    "Section Header": "Corp Divider",
    "Title Only": "Corp Content",
}
WIDE_WIDTH_EMU = 12192000  # 13.333 in x 7.5 in = 16:9
ALT_LABEL = "Alt classification"


def make_alt_template(folder: Path) -> Path:
    from pptx import Presentation

    prs = Presentation()
    prs.slide_width = WIDE_WIDTH_EMU
    for layout in prs.slide_layouts:
        if layout.name in RENAMES:
            layout._element.cSld.name = RENAMES[layout.name]
    id_list = prs.slide_master.slide_layouts._sldLayoutIdLst
    items = list(id_list)
    for item in items:
        id_list.remove(item)
    for item in items[3:] + items[:3]:  # rotate: every layout index changes
        id_list.append(item)
    path = folder / "alt.pptx"
    prs.save(str(path))
    return path


def alt_map_data() -> dict:
    content = {"layout_name": "Corp Content", "fallback_index": 0, "placeholders": {"title": 0}}
    box = [0.6, 1.6, 12.1, 5.2]
    layouts = {
        "title": {"layout_name": "Corp Cover", "fallback_index": 1, "placeholders": {"title": 0, "subtitle": 1}},
        "section": {"layout_name": "Corp Divider", "fallback_index": 1, "placeholders": {"title": 0, "body": 1}},
        "narrative": {"layout_name": "Corp Text", "fallback_index": 1, "placeholders": {"title": 0, "body": 1}},
    }
    rows = {"table": 12, "findings": 8, "attention_list": 12, "provenance": 18}
    for kind in SLIDE_KINDS:
        if kind not in layouts:
            layouts[kind] = {**content, "content_box_in": box, "max_rows": rows.get(kind, 12)}
    return {
        "name": "alt",
        "template": "alt.pptx",
        "slide_size": "16:9",
        "fonts": {"title": "Arial", "body": "Arial", "body_pt": 12, "table_pt": 9},
        "series_colors": ["#004B87", "#6CACE4", "#E87722", "#5E6A71"],
        "footer_text": "{period} weekly | {report_title}",
        "classification_label": ALT_LABEL,
        "layouts": layouts,
    }


@pytest.fixture
def alt_map(tmp_path: Path) -> Path:
    folder = tmp_path / "alt-map"
    folder.mkdir()
    make_alt_template(folder)
    path = folder / "alt.map.yaml"
    path.write_text(yaml.safe_dump(alt_map_data(), sort_keys=False), encoding="utf-8")
    return path


def test_alt_map_resolves_layouts_by_name_and_proofs(tmp_path, alt_map):
    loaded = load_template_map(str(alt_map))
    assert loaded.map.name == "alt" and loaded.template_path == (alt_map.parent / "alt.pptx").resolve()
    result = template_proof(loaded, tmp_path / "proof", data_class="synthetic")
    assert result["fallbacks"] == [] and len(result["kinds"]) == 11
    assert {s["layout"] for s in result["slides"]} == set(RENAMES.values())
    assert deck_problems(Path(result["path"])) == []
    assert open_deck(Path(result["path"])).slide_width == WIDE_WIDTH_EMU


def _titles_in_title_placeholders(path: Path) -> list[str]:
    from pptx.enum.shapes import PP_PLACEHOLDER

    titles = []
    for n, slide in enumerate(open_deck(path).slides, start=1):
        title = slide.shapes.title
        assert title is not None, f"slide {n} has no title placeholder"
        assert title.is_placeholder and title.placeholder_format.type in {
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        }
        assert title.text_frame.text.strip(), f"slide {n} has an empty title"
        titles.append(title.text_frame.text)
    return titles


def test_weekly_deck_builds_with_the_alt_map(ops_profile_rw, alt_map):
    paths = ops_profile_rw.paths
    result = build_report(paths, "weekly", "2026-W35", ["pptx", "xlsx", "md"], "none", template_map=str(alt_map))
    by_format = {a["format"]: a for a in result["artifacts"]}
    loaded = load_template_map(str(alt_map), paths)
    assert by_format["pptx"]["template_map"] == {"name": "alt", "sha256": loaded.sha256}
    assert by_format["xlsx"]["template_map"] is None and Path(by_format["md"]["path"]).is_file()
    deck = Path(by_format["pptx"]["path"])
    assert deck.name == "weekly_2026-W35_SYNTHETIC.pptx"
    assert deck_problems(deck) == []
    alt_titles = _titles_in_title_placeholders(deck)
    prs = open_deck(deck)
    assert prs.slide_width == WIDE_WIDTH_EMU
    assert {s.slide_layout.name for s in prs.slides} <= set(RENAMES.values())
    for slide in prs.slides:
        text = slide_text(slide, notes=False)
        assert ALT_LABEL in text and "2026-W35 weekly | Weekly Application Operations Review" in text
    conn = sqlite3.connect(str(paths.db))
    try:
        stored = conn.execute("SELECT template_map_sha FROM report_artifact WHERE format = 'pptx'").fetchall()
    finally:
        conn.close()
    assert stored == [(loaded.sha256,)]

    neutral = build_report(paths, "weekly", "2026-W35", ["pptx"], "none")
    neutral_deck = Path(neutral["artifacts"][0]["path"])
    assert neutral["artifacts"][0]["template_map"]["name"] == "neutral"
    assert _titles_in_title_placeholders(neutral_deck) == alt_titles, "a map change must not change the deck content"


def test_cli_build_honours_template_map(ops_profile_rw, alt_map):
    from sed.cli import app

    args = ["report", "build", "weekly", "--period", "2026-W35", "--format", "pptx", "--ai", "none"]
    args += ["--template-map", str(alt_map), "--data-dir", str(ops_profile_rw.paths.data_dir), "--json"]
    result = CliRunner().invoke(app, args)
    assert result.exit_code == 0, result.stdout
    payload = json.loads(result.stdout.strip().splitlines()[-1])
    assert payload["artifacts"][0]["template_map"]["name"] == "alt"
    bad = CliRunner().invoke(app, [*args[:-3], "--template-map", "missing-map", *args[-3:-1], "--json"])
    assert bad.exit_code == 2
