"""PPTX engine on hand-built snapshots: every kind, pagination, placeholders, bounds, notes, provenance, AI modes.

`deck_problems`, `slide_text` and `hand_snapshot` are shared with the other deck tests (imported, not fixtures).
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

import pytest

from sed.errors import ValidationFailed
from sed.reports.build import artifact_name
from sed.reports.pptx_builder import (
    DRAFT_STAMP,
    EXCLUDED_TEXT,
    NO_ROWS_TEXT,
    SYNTHETIC_BANNER,
    build_pptx,
    effective_slides,
    format_value,
    render_deck,
)
from sed.reports.snapshot import Snapshot, fact, format_provenance_line, table
from sed.reports.specs import ChartSpec, KpiSpec, ReportSpec, SlideSpec
from sed.reports.template_map import SLIDE_KINDS, load_template_map

AI_FACT_LABEL = "Injected AI accuracy fact"
AI_FACT_VALUE = 87.3
SUPPRESSED_TITLE = "Injected acknowledged renewal finding"


# -- shared helpers ------------------------------------------------------------------------------------------------


def open_deck(path: Path) -> Any:
    from pptx import Presentation

    return Presentation(str(path))


def slide_text(slide: Any, *, notes: bool = True) -> str:
    """All visible text of a slide (text frames, table cells, chart titles) plus its notes."""
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            parts += [cell.text for row in shape.table.rows for cell in row.cells]
        if getattr(shape, "has_chart", False) and shape.has_chart and shape.chart.has_title:
            parts.append(shape.chart.chart_title.text_frame.text)
    if notes and slide.has_notes_slide:
        parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def deck_problems(path: Path) -> list[str]:
    """Structural rules every deck must satisfy: filled placeholders, no '{{', shapes inside the slide, notes."""
    prs = open_deck(path)
    width, height = prs.slide_width, prs.slide_height
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        for ph in slide.placeholders:
            if not (ph.has_text_frame and ph.text_frame.text.strip()):
                problems.append(f"slide {n}: empty placeholder idx {ph.placeholder_format.idx}")
        for shape in slide.shapes:
            left, top, w, h = shape.left, shape.top, shape.width, shape.height
            if None in (left, top, w, h) or left < 0 or top < 0 or left + w > width or top + h > height:
                problems.append(f"slide {n}: shape '{shape.name}' outside the slide")
        if "{{" in slide_text(slide):
            problems.append(f"slide {n}: unresolved '{{{{' token")
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            problems.append(f"slide {n}: no speaker notes")
    return problems


def ai_run(**overrides: Any) -> dict[str, Any]:
    run = {
        "run_id": "20260901T100000-triage-batch-test",
        "skill": "sed-triage-batch",
        "skill_hash": "abcdef012345" * 3,
        "status": "approved",
        "model_reported": "model-x",
        "reviewed_by": "reviewer-one",
        "reviewed_at": "2026-09-02T08:00:00Z",
        "sample_n": 30,
        "sample_accuracy": 0.9,
        "sample_ci_low": 0.74,
        "sample_ci_high": 0.97,
        "used_for": ["categories"],
    }
    return {**run, **overrides}


def rows(n: int, **extra: Any) -> list[dict[str, Any]]:
    return [{"name": f"Row {i}", "value": i * 10, "share": i / 100, **extra} for i in range(n)]


def hand_snapshot(**overrides: Any) -> Snapshot:
    columns = [("name", "Name", "text"), ("value", "Value", "count"), ("share", "Share", "ratio")]
    base: dict[str, Any] = {
        "snapshot_id": "snap-weekly-2026-W35-handbuilt",
        "report_key": "weekly",
        "period": "2026-W35",
        "vendor_id": None,
        "as_of": "2026-08-31",
        "data_class": "synthetic",
        "reporting_tz": "Europe/Paris",
        "sla_source": "task_sla",
        "facts": {
            "inc.opened": fact(47, "count", "Incidents opened", "inc.opened"),
            "inc.opened.avg4w": fact(115.25, "number", "Opened, 4-week average", "inc.opened"),
            "inc.sla.pct": fact(79.07, "pct", "SLA met"),
            "inc.sla.delta": fact(-7.13, "pp", "SLA vs average"),
            "cost.idle": fact(187000.0, "eur", "Idle cost"),
            "mttr": fact(26.24, "hours", "MTTR"),
            "ai.accuracy": fact(AI_FACT_VALUE, "pct", AI_FACT_LABEL),
            **{f"extra.{i}": fact(i, "count", f"Extra {i}") for i in range(8)},
        },
        "tables": {
            "trend": table(
                "Trend",
                [("week", "Week", "text"), ("opened", "Opened", "count"), ("resolved", "Resolved", "count")],
                [{"week": f"W{i}", "opened": 10 + i, "resolved": 9 + i} for i in range(12)],
            ),
            "big": table("Big table", columns, rows(30)),
            "small": table("Small table", columns, rows(3)),
            "empty": table("Empty table", columns, []),
            "findings": table(
                "Findings",
                [("severity", "Severity", "text"), ("kind", "Kind", "text"), ("title", "Finding", "text")],
                [{"severity": "high", "kind": "renewal_risk", "title": f"Finding {i}"} for i in range(17)],
            ),
            "attention": table("Attention", columns, rows(13)),
            "ai_table": table("AI categories", columns, rows(4)),
        },
        "freshness": [
            {
                "mapping_name": "servicenow_incident",
                "last_import": "2026-09-01T00:00:00Z",
                "files": 2,
                "latest_as_of": "2026-09-01",
            }
        ],
        "input_batches": [],
        "sha256": "5" * 64,
        "created_at": "2026-09-01T00:00:00Z",
        "git_commit": None,
        "ai_runs": [ai_run()],
        "ai_derived_tables": ["ai_table"],
        "ai_derived_facts": ["ai.accuracy"],
        "period_end": "2026-08-31",
        "data_as_of": "2026-09-01",
        "suppressed_findings": [
            {
                "stable_key": "k1",
                "kind": "renewal_risk",
                "title": SUPPRESSED_TITLE,
                "status": "acknowledged",
                "suppress_until": None,
            }
        ],
    }
    base.update(overrides)
    return Snapshot(**base)


def all_kinds_spec() -> ReportSpec:
    slides = [
        SlideSpec(kind="title", subtitle="Week {period}"),
        SlideSpec(kind="section", title="Section", subtitle="Subtitle for {report_title}"),
        SlideSpec(
            kind="kpis",
            title="KPIs",
            kpis=[
                KpiSpec(fact="inc.opened", compare="inc.opened.avg4w"),
                KpiSpec(fact="inc.sla.pct", delta="inc.sla.delta"),
                KpiSpec(fact="ai.accuracy"),
                KpiSpec(fact="cost.idle"),
            ],
            notes="Opened {{f:inc.opened}}; AI {{f:ai.accuracy}}",
        ),
        SlideSpec(
            kind="line_chart",
            table="trend",
            chart=ChartSpec(type="line", categories="week", series=["opened", "resolved"]),
        ),
        SlideSpec(
            kind="bar_chart", table="trend", chart=ChartSpec(categories="week", series=["opened"], title="Opened")
        ),
        SlideSpec(
            kind="stacked_bar",
            table="trend",
            chart=ChartSpec(type="bar", categories="week", series=["opened", "resolved"]),
        ),
        SlideSpec(kind="table", table="small"),
        SlideSpec(kind="table", table="ai_table"),
        SlideSpec(kind="narrative", title="Highlights", ai_section_key="highlights"),
        SlideSpec(kind="findings", table="findings"),
        SlideSpec(kind="attention_list", table="attention", columns=["name", "value"]),
        SlideSpec(kind="provenance"),
    ]
    return ReportSpec(report="weekly", title="Hand-built review", kpis=[], sheets=[], slides=slides)


@pytest.fixture(scope="module")
def neutral():
    return load_template_map("neutral")


def build(tmp_path: Path, tmap: Any, spec: ReportSpec, snapshot: Snapshot | None = None, ai_mode: str = "approved"):
    out = tmp_path / f"deck_{ai_mode}.pptx"
    build_pptx(snapshot or hand_snapshot(), spec, tmap, out, ai_mode=ai_mode, generated_at="2026-09-01T12:00:00Z")
    return out, open_deck(out)


def titles(prs: Any) -> list[str]:
    return [s.shapes.title.text if s.shapes.title is not None else "" for s in prs.slides]


# -- tests ---------------------------------------------------------------------------------------------------------


def test_every_kind_renders_with_the_neutral_map(tmp_path, neutral):
    out = tmp_path / "kinds.pptx"
    manifest = render_deck(
        hand_snapshot(),
        all_kinds_spec(),
        neutral,
        out,
        ai_mode="approved",
        generated_at="2026-09-01T12:00:00Z",
        narratives={"highlights": ["First point", "Second point"]},
    )
    assert {s["kind"] for s in manifest["slides"]} == set(SLIDE_KINDS)
    assert not any(s["fallback_used"] for s in manifest["slides"])
    assert deck_problems(out) == []


def test_structure_placeholders_tokens_bounds_and_charts(tmp_path, neutral):
    out, prs = build(tmp_path, neutral, all_kinds_spec())
    assert deck_problems(out) == []
    assert titles(prs)[0] == "Hand-built review"
    chart_slides = [s for s in prs.slides if any(getattr(sh, "has_chart", False) for sh in s.shapes)]
    assert len(chart_slides) == 3
    chart_types = {str(sh.chart.chart_type) for s in chart_slides for sh in s.shapes if getattr(sh, "has_chart", False)}
    assert any("LINE" in t for t in chart_types) and any("STACKED" in t for t in chart_types)
    kpi_notes = next(s.notes_slide.notes_text_frame.text for s in prs.slides if s.shapes.title.text == "KPIs")
    assert "Opened 47" in kpi_notes and "AI 87.3%" in kpi_notes


def test_pagination_counts(tmp_path, neutral):
    spec = ReportSpec(
        report="weekly",
        title="Pagination",
        kpis=[],
        sheets=[],
        slides=[
            SlideSpec(kind="table", title="Big", table="big"),
            SlideSpec(kind="table", title="Big ten", table="big", max_rows=10),
            SlideSpec(kind="findings", title="Findings", table="findings"),
            SlideSpec(kind="attention_list", title="Attention", table="attention"),
            SlideSpec(kind="kpis", title="Tiles", kpis=[KpiSpec(fact=f"extra.{i % 8}") for i in range(14)]),
            SlideSpec(kind="table", title="Skipped", table="empty", when="has_rows"),
            SlideSpec(kind="table", title="Kept empty", table="empty"),
            SlideSpec(kind="provenance", title="Provenance"),
        ],
    )
    out, prs = build(tmp_path, neutral, spec)
    names = titles(prs)
    layouts = neutral.map.layouts
    assert names.count("Big (1/3)") == 1 and "Big (3/3)" in names and math.ceil(30 / layouts["table"].max_rows) == 3
    assert [n for n in names if n.startswith("Big ten")] == ["Big ten (1/3)", "Big ten (2/3)", "Big ten (3/3)"]
    assert len([n for n in names if n.startswith("Findings")]) == math.ceil(17 / layouts["findings"].max_rows)
    assert len([n for n in names if n.startswith("Attention")]) == math.ceil(13 / layouts["attention_list"].max_rows)
    assert len([n for n in names if n.startswith("Tiles")]) == 2
    assert not any(n.startswith("Skipped") for n in names)
    kept = prs.slides[names.index("Kept empty")]
    assert NO_ROWS_TEXT in slide_text(kept, notes=False)
    provenance = [n for n in names if n.startswith("Provenance")]
    assert provenance and names[-len(provenance) :] == provenance
    assert "Skipped: no rows this period" in "\n".join(slide_text(s) for s in prs.slides)
    assert deck_problems(out) == []


def test_notes_carry_snapshot_sources_and_ai_runs(tmp_path, neutral):
    _, prs = build(tmp_path, neutral, all_kinds_spec())
    snap = hand_snapshot()
    line = format_provenance_line(snap.ai_runs[0])
    for slide in prs.slides:
        notes = slide.notes_slide.notes_text_frame.text
        assert snap.snapshot_id in notes and snap.sha256 in notes
        assert "data as of 2026-09-01" in notes and "SLA source: task_sla" in notes and line in notes
    table_notes = next(s.notes_slide.notes_text_frame.text for s in prs.slides if s.shapes.title.text == "Small table")
    assert "Source: snapshot table 'small'" in table_notes and "servicenow_incident" in table_notes


def test_provenance_slide_lists_ai_runs_suppressions_and_omissions(tmp_path, neutral):
    _, prs = build(tmp_path, neutral, all_kinds_spec())
    names = titles(prs)
    assert names[-1].startswith("Provenance")
    text = "\n".join(slide_text(s, notes=False) for s in prs.slides if s.shapes.title.text.startswith("Provenance"))
    run = hand_snapshot().ai_runs[0]
    assert format_provenance_line(run) in text and run["skill_hash"][:12] in text
    assert "Suppressed/acknowledged rule findings" in text and SUPPRESSED_TITLE in text
    assert "Omitted sections" in text and "Highlights" in text and "highlights" in text
    assert "Data as of: 2026-09-01" in text and "period end (exclusive): 2026-08-31" in text
    assert "servicenow_incident" in text and "template map neutral" in text
    assert not any(n == "Highlights" for n in names), "narrative slides are omitted in M2"


def test_synthetic_banner_and_title_slide(tmp_path, neutral):
    _, prs = build(tmp_path, neutral, all_kinds_spec())
    for slide in prs.slides:
        banner = [sh for sh in slide.shapes if sh.name == "sed-synthetic-banner"]
        assert banner and banner[0].text_frame.text == SYNTHETIC_BANNER
        assert any(sh.name == "sed-footer" for sh in slide.shapes)
    assert "SYNTHETIC" in slide_text(prs.slides[0], notes=False)
    assert "Week 2026-W35" in slide_text(prs.slides[0], notes=False)
    _, real = build(tmp_path / "real", neutral, all_kinds_spec(), hand_snapshot(data_class="real"))
    assert not any(sh.name == "sed-synthetic-banner" for s in real.slides for sh in s.shapes)


def test_draft_mode_stamps_every_slide(tmp_path, neutral):
    snap = hand_snapshot()
    out, prs = build(tmp_path, neutral, all_kinds_spec(), snap, ai_mode="draft")
    assert all(
        any(sh.name == "sed-draft-stamp" and sh.text_frame.text == DRAFT_STAMP for sh in s.shapes) for s in prs.slides
    )
    assert artifact_name(snap, "pptx", "draft") == "weekly_2026-W35_SYNTHETIC_DRAFT.pptx"
    assert deck_problems(out) == []
    _, approved = build(tmp_path, neutral, all_kinds_spec(), snap, ai_mode="approved")
    assert not any(sh.name == "sed-draft-stamp" for s in approved.slides for sh in s.shapes)


def test_ai_none_excludes_ai_tables_facts_and_runs(tmp_path, neutral):
    snap = hand_snapshot()
    out, prs = build(tmp_path, neutral, all_kinds_spec(), snap, ai_mode="none")
    assert deck_problems(out) == []
    everything = "\n".join(slide_text(s) for s in prs.slides)
    assert AI_FACT_LABEL not in everything and format_value(AI_FACT_VALUE, "pct") not in everything
    assert format_provenance_line(snap.ai_runs[0]) not in everything
    ai_slide = next(s for s in prs.slides if s.shapes.title.text == "AI categories")
    assert EXCLUDED_TEXT in slide_text(ai_slide, notes=False)
    assert not any(getattr(sh, "has_table", False) and sh.has_table for sh in ai_slide.shapes)
    assert "AI content excluded (--ai none)" in everything and "AI runs: excluded (--ai none)" in everything
    kpi_slide = next(s for s in prs.slides if s.shapes.title.text == "KPIs")
    assert not any(sh.name == "sed-kpi-ai.accuracy" for sh in kpi_slide.shapes)


def test_missing_spec_reference_exits_2(tmp_path, neutral):
    for slide in (
        SlideSpec(kind="table", table="nope"),
        SlideSpec(kind="table", table="small", columns=["name", "missing"]),
        SlideSpec(kind="bar_chart", table="trend", chart=ChartSpec(categories="week", series=["nope"])),
        SlideSpec(kind="kpis", kpis=[KpiSpec(fact="inc.opened", compare="nope.fact")]),
    ):
        spec = ReportSpec(report="weekly", title="Bad", kpis=[], sheets=[], slides=[slide])
        with pytest.raises(ValidationFailed) as exc:
            build(tmp_path, neutral, spec)
        assert exc.value.exit_code == 2


def test_untrusted_cell_text_stays_plain(tmp_path, neutral):
    hostile = "=cmd|' /C calc'!A0 {{f:inc.opened}} \x07bell"
    snap = hand_snapshot()
    snap.tables["small"]["rows"][0]["name"] = hostile
    spec = ReportSpec(report="weekly", title="T", kpis=[], sheets=[], slides=[SlideSpec(kind="table", table="small")])
    _, prs = build(tmp_path, neutral, spec, snap)
    cells = [
        c.text
        for sh in prs.slides[0].shapes
        if getattr(sh, "has_table", False) and sh.has_table
        for row in sh.table.rows
        for c in row.cells
    ]
    assert "=cmd|' /C calc'!A0 {{f:inc.opened}} bell" in cells


def test_default_deck_when_spec_declares_no_slides(tmp_path, neutral):
    from sed.reports.specs import SheetSpec

    spec = ReportSpec(
        report="weekly",
        title="Defaults",
        kpis=[KpiSpec(fact="inc.opened", compare="inc.opened.avg4w"), KpiSpec(fact="not.in.snapshot")],
        sheets=[
            SheetSpec(table="trend", sheet="Trend", chart=ChartSpec(categories="week", series=["opened"])),
            SheetSpec(table="small", sheet="Small"),
            SheetSpec(table="absent", sheet="Absent"),
        ],
    )
    kinds = [s.kind for s in effective_slides(spec, hand_snapshot())]
    assert kinds == ["title", "kpis", "bar_chart", "table", "table", "provenance"]
    out, _ = build(tmp_path, neutral, spec)
    assert deck_problems(out) == []


def test_format_value_units():
    assert format_value(1234, "count") == "1,234"
    assert format_value(79.07, "pct") == "79.1%"
    assert format_value(-7.13, "pp") == "-7.1 pp"
    assert format_value(0.31, "ratio") == "31%"
    assert format_value(26.24, "hours") == "26.2 h"
    assert format_value(187000.0, "eur") == "€187,000"
    assert format_value("2026-08-25T08:23:18Z", "datetime") == "2026-08-25 08:23"
    assert format_value(None, "pct", empty="n/a") == "n/a"
    assert format_value("n/a", "count") == "n/a"
