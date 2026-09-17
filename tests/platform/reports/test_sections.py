"""AI-drafted report sections in builds: approved sections render with snapshot numbers in the deck, workbook and
Markdown; stale (a cited fact changed) and blocked (a cited finding unapproved) sections are left out; draft mode shows
drafts; --require-complete refuses incomplete builds; readiness is read-only."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sed import db
from sed.ai import findings as F
from sed.ai.review import review_findings
from sed.errors import PreconditionFailed, ValidationFailed
from sed.reports import sections
from sed.reports.build import build_report
from sed.reports.snapshot import create_snapshot
from sed.reports.specs import load_report_spec
from tests.platform.reports.test_pptx_builder import open_deck, slide_text

PERIOD = "2026-W35"


def _facts(paths) -> dict:
    conn = db.connect(paths.db)
    try:
        return create_snapshot(conn, paths, "weekly", PERIOD).facts
    finally:
        conn.close()


def _run(conn, run_id: str) -> None:
    conn.execute(
        "INSERT OR IGNORE INTO ai_run (run_id, skill, skill_hash, schema_version, invoked_via, profile, status, "
        "started_at) VALUES (?, 'sed-draft-report', 'h', 1, 'workflow', 'synthetic', 'completed', ?)",
        (run_id, "2026-09-01T07:00:00Z"),
    )


def _section(paths, run_id, key, body, facts, *, cited=(), drafted=None) -> str:
    conn = db.connect(paths.db)
    try:
        with db.write_tx(conn):
            _run(conn, run_id)
            used = {t: facts[t]["value"] for t in sections.tokens(body)}
            return F.upsert_draft(
                conn,
                run_id=run_id,
                stable_key=sections.stable_key("weekly", PERIOD, None, key),
                kind="report_section",
                title=f"weekly {PERIOD}: {key}",
                body_md=body,
                severity=None,
                confidence=0.8,
                payload={
                    "report": "weekly",
                    "period": PERIOD,
                    "section_key": key,
                    "facts": drafted if drafted is not None else used,
                    "cited_finding_ids": list(cited),
                    "slide_headline": None,
                },
                subject_type="report",
                subject_id=sections.subject_id("weekly", PERIOD, None),
                period=PERIOD,
            )
    finally:
        conn.close()


def _cluster(paths) -> str:
    conn = db.connect(paths.db)
    try:
        with db.write_tx(conn):
            _run(conn, "rec-sec")
            return F.upsert_draft(
                conn,
                run_id="rec-sec",
                stable_key="issue_cluster:x",
                kind="issue_cluster",
                title="Cluster",
                body_md="text",
                severity="high",
                confidence=0.8,
                payload={},
            )
    finally:
        conn.close()


def test_approved_section_renders_with_snapshot_numbers(ops_profile_rw):
    paths = ops_profile_rw.paths
    facts = _facts(paths)
    body = "- **{{f:inc.opened}} incidents** opened this week.\n- Backlog at {{f:inc.backlog}}."
    headline = _section(paths, "draft-1", "headline", body, facts)
    review_findings(paths, [headline], "approve", "tester")

    result = build_report(paths, "weekly", PERIOD, ["pptx", "xlsx", "md"], "approved")
    assert result["readiness"]["ai_sections_shown"] == 1 and result["ai_runs"] == ["draft-1"]
    assert any(o.startswith("highlights: no draft") for o in result["readiness"]["omitted"])
    artifacts = {a["format"]: Path(a["path"]) for a in result["artifacts"]}
    opened = f"{facts['inc.opened']['value']:,} incidents opened this week."
    deck = open_deck(artifacts["pptx"])
    slide = next(s for s in deck.slides if s.shapes.title.text_frame.text == "Headline")
    assert opened in slide_text(slide) and "{{f:" not in slide_text(slide) and "**" not in slide_text(slide)
    markdown = artifacts["md"].read_text(encoding="utf-8")
    assert "**Headline**" in markdown and f"**{facts['inc.opened']['value']:,} incidents**" in markdown
    sheet = load_workbook(artifacts["xlsx"], read_only=True)["AI sections"]
    assert any(opened in str(c) for row in sheet.iter_rows(values_only=True) for c in row if c)

    none = build_report(paths, "weekly", PERIOD, ["md", "pptx"], "none")
    assert none["readiness"]["ai_sections_shown"] == 0
    none_md = next(Path(a["path"]) for a in none["artifacts"] if a["format"] == "md").read_text(encoding="utf-8")
    assert "Headline" not in none_md
    conn = sqlite3.connect(str(paths.db))
    try:
        omitted = conn.execute(
            "SELECT unapproved_omitted_json FROM report_artifact WHERE ai_mode = 'none' AND format = 'md'"
        ).fetchone()[0]
    finally:
        conn.close()
    assert "headline: excluded (--ai none)" in json.loads(omitted)


def test_stale_blocked_draft_and_require_complete(ops_profile_rw):
    paths = ops_profile_rw.paths
    facts = _facts(paths)
    spec = load_report_spec("weekly", paths)
    wrong = {"inc.opened": facts["inc.opened"]["value"] + 1}
    stale = _section(paths, "draft-2", "highlights", "{{f:inc.opened}} opened.", facts, drafted=wrong)
    cluster = _cluster(paths)
    blocked = _section(paths, "draft-2", "recurring_issues", "One cluster.", facts, cited=[cluster])
    review_findings(paths, [cluster], "approve", "tester")
    review_findings(paths, [stale, blocked], "approve", "tester")
    review_findings(paths, [cluster], "reject", "tester", note="not a real pattern")  # after the section
    only_draft = _section(paths, "draft-2", "actions", "Call the vendor.", facts)

    conn = db.connect(paths.db)
    try:
        snap = create_snapshot(conn, paths, "weekly", PERIOD)
        states = {s.key: s for s in sections.load_sections(
            conn, spec, "weekly", PERIOD, None, ai_mode="approved", facts=snap.facts
        )}  # fmt: skip
    finally:
        conn.close()
    assert states["highlights"].status == "stale" and "inc.opened changed" in states["highlights"].reasons[0]
    assert states["recurring_issues"].status == "blocked" and "is rejected" in states["recurring_issues"].reasons[0]
    assert states["actions"].status == "draft" and states["headline"].status == "missing"

    with pytest.raises(PreconditionFailed) as refused:
        build_report(paths, "weekly", PERIOD, ["md"], "approved", require_complete=True)
    missing = refused.value.details["sections"]
    assert {m.split(":")[0] for m in missing} == {"headline", "highlights", "lowlights", "recurring_issues", "actions"}
    with pytest.raises(ValidationFailed):
        build_report(paths, "weekly", PERIOD, ["md"], "none", require_complete=True)

    draft = build_report(paths, "weekly", PERIOD, ["md"], "draft")
    shown = draft["readiness"]["ai_sections_shown"]
    text = Path(draft["artifacts"][0]["path"]).read_text(encoding="utf-8")
    assert shown == 2 and "Call the vendor." in text and "One cluster." in text  # drafts ignore citation status
    assert only_draft

    conn = db.connect(paths.db, readonly=True)
    try:
        ready = sections.readiness(conn, paths, "weekly", PERIOD, None)
    finally:
        conn.close()
    by_key = {s["key"]: s for s in ready["sections"]}
    assert ready["sections_required"] == 5 and ready["sections_approved"] == 0 and not ready["complete"]
    assert by_key["recurring_issues"]["approved_status"] == "blocked" and ready["cited_findings_unapproved"] == []
    assert by_key["highlights"]["approved_status"] == "stale" and by_key["actions"]["draft_status"] == "draft"


def test_paragraphs_and_tokens():
    text = "# Title\nFirst line\ncontinues **bold**.\n\n- one [link](#/x)\n1. two\n"
    assert sections.paragraphs(text) == ["Title First line continues bold.", "one link", "two"]
    assert sections.tokens("{{f:a.b}} and {{f:a.b}} {{f:c}}") == ["a.b", "c"]
    assert sections.same_value(12, 12.0) and not sections.same_value(12, 13) and not sections.same_value(True, 1)
