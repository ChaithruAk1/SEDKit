"""`sed report template-inspect` and `template-proof`: JSON shape, starter map, one slide per kind, exit codes."""

from __future__ import annotations

import json
from pathlib import Path

import yaml
from typer.testing import CliRunner

from sed.cli import app
from sed.reports.template_map import SLIDE_KINDS, TemplateMap, default_template_path, load_template_map
from sed.reports.template_tools import inspect_template, template_proof
from tests.platform.reports.test_pptx_builder import deck_problems, open_deck, slide_text

runner = CliRunner()


def invoke_json(args: list[str]) -> tuple[int, dict]:
    result = runner.invoke(app, [*args, "--json"])
    lines = [ln for ln in result.stdout.strip().splitlines() if ln.strip()]
    assert len(lines) == 1, f"expected one JSON line, got {result.stdout!r}"
    return result.exit_code, json.loads(lines[0])


def default_copy(tmp_path: Path) -> Path:
    from pptx import Presentation

    path = tmp_path / "default copy.pptx"
    Presentation().save(str(path))
    return path


def test_inspect_json_shape_on_default_template(tmp_path):
    code, out = invoke_json(["report", "template-inspect", str(default_copy(tmp_path))])
    assert code == 0 and out["ok"] is True
    assert out["slide_size"] == {"width_in": 10.0, "height_in": 7.5, "ratio": "4:3"}
    assert out["masters"] == 1 and out["slides"] == 0 and out["fallback_index_range"] == [0, 10]
    names = [layout["name"] for layout in out["layouts"]]
    assert names[:3] == ["Title Slide", "Title and Content", "Section Header"] and names[5:7] == ["Title Only", "Blank"]
    title = out["layouts"][0]
    assert set(title) == {"index", "master", "position", "name", "placeholders"}
    first = title["placeholders"][0]
    assert set(first) == {"idx", "type", "name", "left_in", "top_in", "width_in", "height_in"}
    assert (first["idx"], first["type"]) == (0, "CENTER_TITLE")
    assert {p["type"] for p in out["layouts"][5]["placeholders"]} >= {"TITLE", "FOOTER"}


def test_suggested_map_is_a_valid_starting_point(tmp_path):
    template = default_copy(tmp_path)
    suggested = inspect_template(template)["suggested_map"]
    TemplateMap.model_validate(suggested)
    suggested["template"] = template.name
    map_path = tmp_path / "starter.map.yaml"
    map_path.write_text(yaml.safe_dump(suggested, sort_keys=False), encoding="utf-8")
    loaded = load_template_map(str(map_path))
    assert loaded.map.layouts["title"].layout_name == "Title Slide"
    assert loaded.map.layouts["table"].layout_name == "Title Only"
    assert loaded.map.layouts["narrative"].placeholders == {"title": 0, "body": 1}


def test_inspect_rejects_potx_and_unreadable_files(tmp_path):
    potx = tmp_path / "corporate.potx"
    potx.write_bytes(b"placeholder")
    code, out = invoke_json(["report", "template-inspect", str(potx)])
    assert code == 2 and "open it in PowerPoint and save as .pptx" in out["error"]["message"]
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a zip archive")
    code, out = invoke_json(["report", "template-inspect", str(broken)])
    assert code == 2 and out["error"]["kind"] == "validation"
    code, _ = invoke_json(["report", "template-inspect", str(tmp_path / "missing.pptx")])
    assert code == 2


def test_template_proof_renders_every_kind(tmp_path):
    result = template_proof(load_template_map("neutral"), tmp_path / "proof", data_class="synthetic")
    path = Path(result["path"])
    assert path.name == "template-proof_neutral_SYNTHETIC.pptx" and path.is_file()
    assert result["kinds"] == sorted(SLIDE_KINDS) and len(result["kinds"]) == 11
    assert result["slides"][0]["kind"] == "title" and result["slides"][-1]["kind"] == "provenance"
    assert result["fallbacks"] == [] and result["template"] is None
    assert deck_problems(path) == []
    prs = open_deck(path)
    assert len(prs.slides) == len(result["slides"])
    narrative = next(s for s, m in zip(prs.slides, result["slides"], strict=True) if m["kind"] == "narrative")
    assert "First example bullet" in slide_text(narrative, notes=False)


def test_template_proof_cli_uses_profile_out_dir_and_map(tmp_path, data_root):
    code, out = invoke_json(["report", "template-proof", "--map", "neutral"])
    assert code == 0 and out["map"] == "neutral" and len(out["kinds"]) == 11
    assert Path(out["path"]).parent == data_root / "synthetic" / "out" / "template-proof"
    code, out = invoke_json(["report", "template-proof", "--out", str(tmp_path / "elsewhere")])
    assert code == 0 and Path(out["path"]).parent == tmp_path / "elsewhere"
    code, out = invoke_json(["report", "template-proof", "--map", str(tmp_path / "nope.map.yaml")])
    assert code == 2


def test_default_template_path_exists():
    assert default_template_path().is_file()
