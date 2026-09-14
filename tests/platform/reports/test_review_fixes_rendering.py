"""Rendering rules from the M2 review: DRAFT on every sheet, base-currency symbols, fact tokens in slide text,
filename-safe vendor ids, vendor-scoped provenance and reuse of a stored snapshot."""

from __future__ import annotations

import dataclasses
import json
import sqlite3

import pytest

from sed import db
from sed.errors import ValidationFailed
from sed.reports.build import artifact_name, build_report
from sed.reports.pptx_builder import check_spec_refs, format_value
from sed.reports.snapshot import create_snapshot, money_num_format
from sed.reports.specs import ReportSpec, SheetSpec, SlideSpec
from sed.reports.xlsx_builder import DRAFT_BANNER, build_xlsx
from tests.platform.reports.test_pptx_builder import hand_snapshot


def test_draft_workbooks_are_stamped_on_every_sheet(tmp_path):
    from openpyxl import load_workbook

    spec = ReportSpec(report="weekly", title="Draft", kpis=[], sheets=[SheetSpec(table="small", sheet="Small")])
    out = tmp_path / "draft.xlsx"
    build_xlsx(hand_snapshot(), spec, out, ai_mode="draft", generated_at="2026-09-01T00:00:00Z")
    book = load_workbook(out)
    assert book.sheetnames == ["Summary", "Small", "Definitions", "Provenance"]
    for sheet in book.worksheets:
        assert sheet.cell(row=1, column=1).value == DRAFT_BANNER, sheet.title
    approved = tmp_path / "approved.xlsx"
    build_xlsx(hand_snapshot(), spec, approved, ai_mode="approved", generated_at="2026-09-01T00:00:00Z")
    assert all(s.cell(row=1, column=1).value != DRAFT_BANNER for s in load_workbook(approved).worksheets)


def test_colour_rules_skip_blank_cells(tmp_path):
    from openpyxl import load_workbook

    from sed.reports.specs import ConditionalSpec

    snap = hand_snapshot()
    rows = snap.tables["small"]["rows"]
    snap.tables["small"]["rows"] = [{**rows[0], "value": None}, *rows[1:]]
    sheet = SheetSpec(table="small", sheet="Small", conditional=[ConditionalSpec(column="value", rule="<", value=90)])
    out = tmp_path / "blanks.xlsx"
    build_xlsx(
        snap, ReportSpec(report="weekly", title="C", kpis=[], sheets=[sheet]), out, ai_mode="none", generated_at="x"
    )
    rules = [rule for cf in load_workbook(out)["Small"].conditional_formatting for rule in cf.rules]
    assert [(r.type, r.stopIfTrue) for r in sorted(rules, key=lambda r: r.priority)][:2] == [
        ("containsBlanks", True),
        ("cellIs", None),
    ]


def test_draft_markdown_is_stamped(ops_profile_rw):
    result = build_report(ops_profile_rw.paths, "weekly", "2026-W35", ["md"], "draft")
    with open(result["artifacts"][0]["path"], encoding="utf-8") as fh:
        assert fh.read().startswith("> **DRAFT**")


def test_money_uses_the_base_currency(tmp_path):
    from openpyxl import load_workbook

    assert format_value(45250.0, "eur") == "€45,250"
    assert format_value(45250.0, "eur", currency="USD") == "$45,250"
    assert format_value(45250.0, "eur", currency="CHF") == "CHF 45,250"
    assert (money_num_format("EUR"), money_num_format("CHF")) == ("€#,##0", '"CHF "#,##0')
    snap = dataclasses.replace(hand_snapshot(), base_currency="CHF")
    spec = ReportSpec(report="weekly", title="Money", kpis=[], sheets=[])
    out = tmp_path / "money.xlsx"
    build_xlsx(snap, spec, out, ai_mode="none", generated_at="2026-09-01T00:00:00Z")
    rows = load_workbook(out)["Provenance"].iter_rows(values_only=True)
    provenance = {row[0]: row[1] for row in rows if row and row[0]}
    assert provenance["Base currency"] == "CHF"


def test_fact_tokens_in_slide_text_must_exist_and_be_well_formed():
    snap = hand_snapshot()
    good = ReportSpec(
        report="weekly", title="T", kpis=[], sheets=[], slides=[SlideSpec(kind="section", notes="{{f:inc.opened}}")]
    )
    check_spec_refs(snap, good)
    for slide in (
        SlideSpec(kind="section", notes="SLA source: {{f:inc.sla.sorce}}"),
        SlideSpec(kind="section", title="Opened {{f:inc.opened}"),
    ):
        spec = ReportSpec(report="weekly", title="T", kpis=[], sheets=[], slides=[slide])
        with pytest.raises(ValidationFailed) as err:
            check_spec_refs(snap, spec)
        assert err.value.details


@pytest.mark.parametrize(
    ("vendor_id", "segment"),
    [("V001", "V001"), ("SUP/0042", "SUP_0042"), ("../../x", "_.._x"), ("A:B*C", "A_B_C")],
)
def test_vendor_ids_become_safe_filename_segments(vendor_id, segment):
    snap = dataclasses.replace(hand_snapshot(), report_key="vendor", period="2026-Q3", vendor_id=vendor_id)
    name = artifact_name(snap, "xlsx", "none")
    assert name == f"vendor_2026-Q3_{segment}_SYNTHETIC.xlsx"
    assert "/" not in name and "\\" not in name


def test_vendor_provenance_lists_only_that_vendors_suppressed_findings(ops_profile_rw):
    paths = ops_profile_rw.paths
    vendor = ops_profile_rw.ids["vendor_p2"]
    conn = db.connect(paths.db)
    try:
        create_snapshot(conn, paths, "weekly", "2026-W35")  # refreshes the rule findings
        other = conn.execute(
            "SELECT f.finding_id, f.title FROM finding f JOIN contract c ON c.contract_id = f.subject_id "
            "WHERE f.origin = 'rule' AND f.subject_type = 'contract' AND COALESCE(c.vendor_id, '') != ? LIMIT 1",
            (vendor,),
        ).fetchone()
        assert other, "the fixture has contract findings for other vendors"
        with db.write_tx(conn):
            conn.execute("UPDATE finding SET status = 'acknowledged' WHERE finding_id = ?", (other["finding_id"],))
        vendor_snap = create_snapshot(conn, paths, "vendor", "2026-Q3", vendor)
        weekly_snap = create_snapshot(conn, paths, "weekly", "2026-W35")
    finally:
        conn.close()
    assert other["title"] not in [f["title"] for f in vendor_snap.suppressed_findings]
    assert other["title"] in [f["title"] for f in weekly_snap.suppressed_findings]


def test_a_reused_snapshot_keeps_its_stored_metadata_and_a_new_import_gives_a_new_one(ops_profile_rw):
    paths = ops_profile_rw.paths
    conn = db.connect(paths.db)
    try:
        first = create_snapshot(conn, paths, "weekly", "2026-W35")
        again = create_snapshot(conn, paths, "weekly", "2026-W35")
        assert again.snapshot_id == first.snapshot_id and again.created_at == first.created_at
        with db.write_tx(conn):
            conn.execute(
                "INSERT INTO import_batch (file_name, file_sha256, mapping_name, mapping_sha256, load_mode, status, "
                "rows_read, imported_at) VALUES ('jira_export_NEW.csv', 'sha-new', 'jira_issues', 'm', 'delta', "
                "'completed', 0, '2026-09-02T00:00:00Z')"
            )
        after_import = create_snapshot(conn, paths, "weekly", "2026-W35")
    finally:
        conn.close()
    assert after_import.facts == first.facts and after_import.snapshot_id != first.snapshot_id
    stored = sqlite3.connect(str(paths.db))
    try:
        row = stored.execute(
            "SELECT input_batches_json FROM report_snapshot WHERE snapshot_id = ?", (after_import.snapshot_id,)
        ).fetchone()
    finally:
        stored.close()
    assert json.loads(row[0]) == [b["batch_id"] for b in after_import.input_batches]
