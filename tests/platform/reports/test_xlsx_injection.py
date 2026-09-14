"""Ticket text that looks like a formula or a URL stays a plain string in workbooks (and in decks)."""

from __future__ import annotations

from sed.reports.pptx_builder import build_pptx
from sed.reports.specs import ReportSpec, SheetSpec, SlideSpec
from sed.reports.template_map import load_template_map
from sed.reports.xlsx_builder import build_xlsx
from tests.platform.reports.test_pptx_builder import hand_snapshot, open_deck, slide_text

INJECTIONS = [
    '=HYPERLINK("https://example.invalid/x","click")',
    "+SUM(1,2)",
    "-2+3",
    "@cmd",
    "https://example.invalid/ticket",
]


def _snapshot():
    snap = hand_snapshot()
    rows = snap.tables["small"]["rows"]
    snap.tables["small"]["rows"] = [{**rows[0], "name": text} for text in INJECTIONS]
    return snap


def test_formula_injection_strings_stay_strings_in_xlsx(tmp_path):
    from openpyxl import load_workbook
    from python_calamine import CalamineWorkbook

    snap = _snapshot()
    spec = ReportSpec(report="weekly", title="Injection", kpis=[], sheets=[SheetSpec(table="small", sheet="Small")])
    out = tmp_path / "injection.xlsx"
    build_xlsx(snap, spec, out, ai_mode="none", generated_at="2026-09-01T00:00:00Z")

    values = CalamineWorkbook.from_path(str(out)).get_sheet_by_name("Small").to_python()
    first_column = [row[0] for row in values[3 : 3 + len(INJECTIONS)]]
    assert first_column == INJECTIONS

    sheet = load_workbook(out)["Small"]
    cells = [sheet.cell(row=4 + i, column=1) for i in range(len(INJECTIONS))]
    assert [c.value for c in cells] == INJECTIONS
    assert {c.data_type for c in cells} == {"s"}
    assert not any(c.hyperlink for c in cells)


def test_formula_injection_strings_are_plain_text_in_pptx(tmp_path):
    spec = ReportSpec(
        report="weekly", title="Injection", kpis=[], sheets=[], slides=[SlideSpec(kind="table", table="small")]
    )
    out = tmp_path / "injection.pptx"
    build_pptx(
        _snapshot(), spec, load_template_map("neutral"), out, ai_mode="none", generated_at="2026-09-01T00:00:00Z"
    )
    text = slide_text(open_deck(out).slides[0], notes=False)
    for value in INJECTIONS:
        assert value in text


def test_formula_injection_strings_stay_strings_in_chart_workbooks(tmp_path):
    """python-pptx embeds each chart's data in a workbook; imported category names must not become formulas there."""
    import re
    import zipfile

    from sed.reports.specs import ChartSpec

    spec = ReportSpec(
        report="weekly",
        title="Injection",
        kpis=[],
        sheets=[],
        slides=[SlideSpec(kind="bar_chart", table="small", chart=ChartSpec(categories="name", series=["value"]))],
    )
    out = tmp_path / "chart_injection.pptx"
    build_pptx(
        _snapshot(), spec, load_template_map("neutral"), out, ai_mode="none", generated_at="2026-09-01T00:00:00Z"
    )
    with zipfile.ZipFile(out) as deck:
        embedded = [name for name in deck.namelist() if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")]
        assert embedded
        for name in embedded:
            with zipfile.ZipFile(deck.open(name)) as book:
                xml = "".join(book.read(n).decode("utf-8") for n in book.namelist() if n.endswith(".xml"))
                assert not re.search(r"<f[ >]", xml), name
                assert not re.search(r"<hyperlink[ >/]", xml), name
                assert "<t>=HYPERLINK(" in xml, "the category survives as plain text"
