"""P13 leak test across the M2 surfaces: AI packets, API JSON, PPTX and XLSX artifacts.

The generator injects fake names, emails and phone numbers into ticket text (ground_truth/pii_injections.json). None of
them may leave the import pipeline: not in the packets an agent reads, not in API responses, not in built reports.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from python_calamine import CalamineWorkbook

from sed.ai.contract import StartParams
from sed.ai.runs import start_run
from sed.reports.build import build_report
from tests.fake_agent.flow import query
from tests.fixtures.api import api_client


def _injections(profile) -> list[str]:
    values = json.loads((profile.ground_truth / "pii_injections.json").read_text(encoding="utf-8"))
    assert len(values) > 500
    return values


def _leaks(text: str, needles: list[str]) -> list[str]:
    return [n for n in needles if n in text]


def _p13_ticket_ids(profile) -> list[str]:
    with (profile.ground_truth / "ticket_truth.csv").open(encoding="utf-8", newline="") as fh:
        return [f"{r['kind']}:{r['number']}" for r in csv.DictReader(fh) if r["pattern"] == "P13"]


def _pptx_text(path: Path) -> str:
    parts = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                parts += [cell.text for row in shape.table.rows for cell in row.cells]
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def _xlsx_text(path: Path) -> str:
    book = CalamineWorkbook.from_path(str(path))
    return "\n".join(
        str(cell) for name in book.sheet_names for row in book.get_sheet_by_name(name).to_python() for cell in row
    )


def test_no_pii_in_ai_packets(ops_profile_rw):
    needles = _injections(ops_profile_rw)
    plan = start_run(
        ops_profile_rw.paths,
        "sed-triage-batch",
        StartParams(scope="since:2025-03-01", limit=6000, max_items=20000, batch_size=500),
    )
    assert plan.plan["items"] > 1000, plan.plan
    selected = {r[0] for r in query(ops_profile_rw.paths, "SELECT item_id FROM ai_batch_item")}
    assert len(selected & set(_p13_ticket_ids(ops_profile_rw))) >= 50, "packets must include PII-injected tickets"
    run_dir = Path(plan.run_dir)
    files = [p for p in (run_dir / "in").rglob("*") if p.is_file()] + [run_dir / "manifest.json"]
    leaks = {p.name: found for p in files if (found := _leaks(p.read_text(encoding="utf-8"), needles))}
    assert leaks == {}


def test_no_pii_in_api_json(ops_profile):
    needles = _injections(ops_profile)
    client = api_client(ops_profile.paths)
    urls = [
        "/api/meta", "/api/findings?status=all", "/api/imports?limit=1000", "/api/dq/unmapped", "/api/runs",
        "/api/ops/overview", "/api/ops/attention?limit=1000", "/api/ops/tickets?page_size=200", "/api/ops/apps",
        "/api/ops/costs", "/api/ops/contracts/renewals?days=1000", "/api/ops/licenses/utilization",
        "/api/ops/tickets?q=regards&page_size=200", "/api/ops/tickets?q=tel&page_size=200",
    ]  # fmt: skip
    urls += [f"/api/ops/tickets/{tid}" for tid in _p13_ticket_ids(ops_profile)[:60]]
    leaks = {}
    details = 0
    for url in urls:
        response = client.get(url)
        assert response.status_code in {200, 404}, (url, response.status_code)
        details += url.startswith("/api/ops/tickets/") and response.status_code == 200
        found = _leaks(response.text, needles)
        if found:
            leaks[url] = found[:3]
    assert details >= 50, "ticket detail responses for PII-injected tickets must be checked"
    assert leaks == {}


def test_no_pii_in_built_reports(ops_profile_rw):
    needles = _injections(ops_profile_rw)
    jobs = [("weekly", "2026-W35", None), ("monthly", "2026-08", None), ("quarterly", "2026-Q3", None),
            ("vendor", "2026-Q3", ops_profile_rw.ids["vendor_p2"])]  # fmt: skip
    leaks = {}
    for report, period, vendor in jobs:
        result = build_report(ops_profile_rw.paths, report, period, ["xlsx", "pptx"], "draft", vendor)
        for artifact in result["artifacts"]:
            path = Path(artifact["path"])
            text = _pptx_text(path) if path.suffix == ".pptx" else _xlsx_text(path)
            found = _leaks(text, needles)
            if found:
                leaks[path.name] = found[:3]
    assert leaks == {}
