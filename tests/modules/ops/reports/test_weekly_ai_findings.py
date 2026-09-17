"""Approved AI findings in the weekly review: only published AI findings of the ops kinds, AI-derived (dropped by
--ai none), with their runs in the provenance, in the workbook sheet, the deck and the Markdown summary."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from sed import db
from sed.ai import findings as F
from sed.ai.review import review_findings
from sed.reports.build import build_report
from sed.reports.pptx_builder import EXCLUDED_TEXT
from tests.modules.ops.reports.test_decks import stored_snapshot
from tests.platform.reports.test_pptx_builder import deck_problems, open_deck, slide_text

RUN = "rec-weekly"


def _findings(paths) -> dict[str, str]:
    conn = db.connect(paths.db)
    try:
        with db.write_tx(conn):
            conn.execute(
                "INSERT INTO ai_run (run_id, skill, skill_hash, schema_version, invoked_via, profile, status, "
                "started_at) VALUES (?, 'sed-find-recurring', 'h', 1, 'workflow', 'synthetic', 'completed', "
                "'2026-09-01T06:00:00Z')",
                (RUN,),
            )

            def draft(key, title, severity, kind="issue_cluster"):
                return F.upsert_draft(
                    conn,
                    run_id=RUN,
                    stable_key=key,
                    kind=kind,
                    title=title,
                    body_md="{{f:T001.tickets}} incidents.",
                    severity=severity,
                    confidence=0.8,
                    payload={"ticket_count": 40, "recommendation": "raise_problem", "evidence": []},
                    subject_type="app",
                    subject_id="APM1001000",
                )

            ids = {
                "approved": draft("k:approved", "Interface timeouts after a change", "high"),
                "critical": draft("k:critical", "Month-end batch failures", "critical"),
                "draft": draft("k:draft", "Unreviewed cluster", "high"),
                "rejected": draft("k:rejected", "Rejected cluster", "medium"),
                "section": draft("k:section", "A report section", "low", kind="report_section"),
            }
    finally:
        conn.close()
    review_findings(paths, [ids["approved"], ids["critical"], ids["section"]], "approve", "tester")
    review_findings(paths, [ids["rejected"]], "reject", "tester", note="noise")
    return ids


def test_weekly_shows_approved_ai_findings_and_none_mode_drops_them(ops_profile_rw):
    paths = ops_profile_rw.paths
    _findings(paths)

    approved = build_report(paths, "weekly", "2026-W35", ["xlsx", "md", "pptx"], "approved")
    snapshot = stored_snapshot(paths, approved["snapshot_id"])
    rows = snapshot["tables"]["ai_findings"]["rows"]
    assert [r["title"] for r in rows] == ["Month-end batch failures", "Interface timeouts after a change"]
    assert rows[0] | {"reviewed_by": "tester", "run_id": RUN, "tickets": 40} == rows[0]
    assert snapshot["facts"]["ai.findings.approved.count"]["value"] == 2
    provenance = snapshot["provenance"]
    assert "ai_findings" in provenance["ai_derived_tables"]
    assert "ai.findings.approved.count" in provenance["ai_derived_facts"]
    assert approved["ai_runs"] == [RUN]
    assert [(r["run_id"], r["used_for"]) for r in provenance["ai_runs"]] == [(RUN, ["ai_findings"])]

    artifacts = {a["format"]: Path(a["path"]) for a in approved["artifacts"]}
    sheet = load_workbook(artifacts["xlsx"], read_only=True)["AI findings"]
    cells = {str(c) for row in sheet.iter_rows(values_only=True) for c in row if c is not None}
    assert {"Month-end batch failures", "Interface timeouts after a change"} <= cells
    assert "Unreviewed cluster" not in cells and "A report section" not in cells
    markdown = artifacts["md"].read_text(encoding="utf-8")
    assert "**AI-assisted findings (approved):**" in markdown and "[critical] Month-end batch failures" in markdown
    deck = open_deck(artifacts["pptx"])
    assert deck_problems(artifacts["pptx"]) == []
    slide = next(s for s in deck.slides if s.shapes.title.text_frame.text.startswith("AI-assisted findings"))
    assert "Interface timeouts after a change" in slide_text(slide)

    none = build_report(paths, "weekly", "2026-W35", ["md", "pptx"], "none")
    none_artifacts = {a["format"]: Path(a["path"]) for a in none["artifacts"]}
    assert none["ai_runs"] == []
    assert "AI-assisted findings" not in none_artifacts["md"].read_text(encoding="utf-8")
    none_deck = open_deck(none_artifacts["pptx"])
    excluded = [s for s in none_deck.slides if s.shapes.title.text_frame.text.startswith("AI-assisted findings")]
    assert excluded and all(EXCLUDED_TEXT in slide_text(s) for s in excluded)
    assert "Month-end batch failures" not in "\n".join(slide_text(s) for s in none_deck.slides)


def test_weekly_without_ai_findings_has_no_ai_findings_slide(ops_profile_rw):
    paths = ops_profile_rw.paths
    result = build_report(paths, "weekly", "2026-W35", ["pptx"], "approved")
    snapshot = stored_snapshot(paths, result["snapshot_id"])
    assert snapshot["tables"]["ai_findings"]["rows"] == []
    deck = open_deck(Path(result["artifacts"][0]["path"]))
    assert not any(s.shapes.title.text_frame.text.startswith("AI-assisted findings") for s in deck.slides)
