"""Report specs for monthly, quarterly and vendor: validation, references into the fixture snapshots, definitions,
workbook builds with --ai none and the ops.config_valid doctor check."""

from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Any

import openpyxl

from sed import db
from sed.reports.snapshot import Snapshot, create_snapshot
from sed.reports.specs import ReportSpec, load_report_spec

REPORTS = {"monthly": ("2026-08", False), "quarterly": ("2026-Q3", False), "vendor": ("2026-Q3", True)}
NARRATIVE_SLOTS = {"monthly": "asks", "quarterly": "decisions_needed", "vendor": "negotiation_points"}
RESERVED_SHEETS = {"Summary", "Definitions", "Provenance"}
# The contract (build spec section 6): fact and table keys each report must carry.
REQUIRED = {
    "monthly": (
        {"inc.opened", "inc.resolved", "inc.sla.pct", "inc.sla.delta_pp_vs_prev_month", "inc.mttr.median_h",
         "inc.backlog", "inc.p1p2.opened", "chg.count", "chg.success.pct", "work.resolved.count"},
        {"trend_6m_by_family", "top_recurring", "problems", "improvements", "upcoming_changes", "findings"},
    ),
    "quarterly": (
        {"cost.actual.qtd", "cost.budget.qtd", "cost.variance.qtd_pct", "cost.actual.ytd", "cost.budget.ytd",
         "license.idle_cost", "renewals.2q.count", "notice.2q.count", "apps.quiet.count"},
        {"spend_by_app", "spend_by_category", "license_idle", "renewals_2q", "risks", "portfolio_health",
         "quiet_apps"},
    ),
    "vendor": (
        {"vendor.name", "vendor.contracts.count", "vendor.contracts.annual_value", "vendor.spend.period",
         "vendor.spend.prev_period", "vendor.sla.pct", "vendor.sla.delta_pp", "vendor.mttr.median_h",
         "vendor.reassign.avg", "vendor.incidents.count", "vendor.license.idle_cost"},
        {"contracts", "spend_trend", "licenses", "sla_trend", "incidents_by_app", "renewal_timeline", "risks"},
    ),
}  # fmt: skip


def _snapshots(profile: Any) -> dict[str, Snapshot]:
    conn = db.connect(profile.paths.db)
    try:
        return {
            key: create_snapshot(conn, profile.paths, key, period, profile.ids["vendor_p2"] if vendor else None)
            for key, (period, vendor) in REPORTS.items()
        }
    finally:
        conn.close()


def _column_keys(snapshot: Snapshot, table: str) -> set[str]:
    return {c["key"] for c in snapshot.tables[table]["columns"]}


def _references(spec: ReportSpec, snap: Snapshot) -> list[str]:
    """Every spec reference that does not resolve in the snapshot."""
    missing: list[str] = []
    kpis = list(spec.kpis) + [k for s in spec.slides for k in (s.kpis or [])]
    for kpi in kpis:
        missing += [f"fact {k}" for k in (kpi.fact, kpi.compare, kpi.delta) if k and k not in snap.facts]
    for sheet in spec.sheets:
        if sheet.table not in snap.tables:
            missing.append(f"sheet table {sheet.table}")
            continue
        keys = _column_keys(snap, sheet.table)
        refs = [c.column for c in sheet.conditional]
        if sheet.chart:
            refs += [sheet.chart.categories, *sheet.chart.series]
        missing += [f"sheet {sheet.sheet} column {c}" for c in refs if c not in keys]
    for slide in spec.slides:
        if slide.table is None:
            continue
        if slide.table not in snap.tables:
            missing.append(f"slide table {slide.table}")
            continue
        keys = _column_keys(snap, slide.table)
        refs = list(slide.columns or [])
        if slide.chart:
            refs += [slide.chart.categories, *slide.chart.series]
        missing += [f"slide '{slide.title}' column {c}" for c in refs if c not in keys]
    return missing


def test_specs_validate_with_sheets_and_slides():
    for key in REPORTS:
        spec = load_report_spec(key, None)
        assert spec.report == key and spec.title and spec.audience and spec.kpis and spec.sheets
        names = [s.sheet for s in spec.sheets]
        assert len(names) == len(set(names)) and not RESERVED_SHEETS & set(names)
        assert all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in names)
        assert spec.slides[0].kind == "title" and spec.slides[-1].kind == "provenance"
        narrative = [s for s in spec.slides if s.kind == "narrative"]
        assert [s.ai_section_key for s in narrative] == [NARRATIVE_SLOTS[key]]
        assert any(s.kind == "kpis" for s in spec.slides)


def test_spec_references_resolve_in_fixture_snapshots(ops_profile_rw):
    from sed.metrics import METRICS
    from sed.modules import metric_definitions
    from sed.modules.ops.ai.definitions import AI_DEFINITIONS
    from sed.modules.ops.reports.definitions import REPORT_DEFINITIONS

    snaps = _snapshots(ops_profile_rw)
    definitions = metric_definitions(ops_profile_rw.paths)
    emitted: dict[str, str] = {}
    for key, snap in snaps.items():
        spec = load_report_spec(key, ops_profile_rw.paths)
        assert _references(spec, snap) == [], key
        facts, tables = REQUIRED[key]
        assert facts <= set(snap.facts) and tables <= set(snap.tables), key
        assert {s.table for s in spec.sheets} == set(snap.tables), f"{key}: every table has a sheet"
        for fact_key, f in snap.facts.items():
            if f["definition"]:
                assert f["definition"] in definitions, f"{key}: {fact_key} -> {f['definition']}"
            if fact_key in REPORT_DEFINITIONS:
                assert f["definition"] == fact_key and f["unit"] == REPORT_DEFINITIONS[fact_key][0]
                emitted[fact_key] = key
    assert set(REPORT_DEFINITIONS) == set(emitted), "every report definition is used by a builder"
    assert not set(REPORT_DEFINITIONS) & (set(METRICS) | set(AI_DEFINITIONS))


def test_workbooks_build_with_ai_none(ops_profile_rw):
    from sed.reports.build import build_report

    paths = ops_profile_rw.paths
    for key, (period, vendor) in REPORTS.items():
        vendor_id = ops_profile_rw.ids["vendor_p2"] if vendor else None
        result = build_report(paths, key, period, ["xlsx"], "none", vendor_id)
        assert result["ai_runs"] == [] and [a["format"] for a in result["artifacts"]] == ["xlsx"]
        target = Path(result["artifacts"][0]["path"])
        assert target.is_file() and target.name.endswith("_SYNTHETIC.xlsx")
        spec = load_report_spec(key, paths)
        wb = openpyxl.load_workbook(target, read_only=True)
        try:
            assert wb.sheetnames == ["Summary", *[s.sheet for s in spec.sheets], "Definitions", "Provenance"]
            defined = {row[0] for row in wb["Definitions"].iter_rows(min_row=2, values_only=True)}
        finally:
            wb.close()
        required_kpis = {k.fact for k in spec.kpis} & REQUIRED[key][0]
        assert required_kpis and required_kpis <= defined, key
    conn = sqlite3.connect(str(paths.db))
    try:
        stored = conn.execute("SELECT DISTINCT ai_mode, ai_run_ids_json FROM report_artifact").fetchall()
    finally:
        conn.close()
    assert stored == [("none", "[]")]


def test_doctor_ops_config_valid_is_ok(ops_profile_rw):
    from sed.modules import doctor_checks

    checks = {c.name: c for c in doctor_checks(ops_profile_rw.paths)}
    assert checks["ops.config_valid"].status == "ok", checks["ops.config_valid"].detail
